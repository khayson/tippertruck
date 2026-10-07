import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from html import escape
import shutil
import os

def create_proposal_with_both_erds():
    backup_file = 'docs/TipperTruck_Proposal_final_year_project.updated.orig_backup.docx'
    output_file = 'docs/TipperTruck_Proposal_final_year_project.updated.final.docx'
    src_file = 'docs/TipperTruck_Proposal_final_year_project.updated.docx'
    
    doc = docx.Document(backup_file)
    body = doc._body._element

    # 1. Update List of Figures
    lof_found = False
    for i, p in enumerate(doc.paragraphs):
        if "Figure 3.5" in p.text and ("tipper_truck_db" in p.text or "Entity-Relationship" in p.text) and "34" in p.text:
            print(f"Found LOF entry at paragraph index {i}: '{p.text}'")
            # Update current paragraph to Figure 3.5a
            p.text = "Figure 3.5a Conceptual Entity-Relationship Model (Peter Chen Notation) ..... 34"
            # Insert Figure 3.5b immediately after
            p_elem = p._element
            new_p_xml = (
                f'<w:p {nsdecls("w")}>'
                f'  <w:pPr><w:spacing w:line="360" w:lineRule="auto"/></w:pPr>'
                f'  <w:r><w:t>Figure 3.5b Physical Relational Database Schema (Crow\'s Foot Notation) ... 35</w:t></w:r>'
                f'</w:p>'
            )
            p_elem.addnext(parse_xml(new_p_xml))
            lof_found = True
            break
            
    if not lof_found:
        print("ERROR: Could not locate LOF Figure 3.5 entry!")
        return

    # 2. Locate Section 3.4.3 and Section 3.4.4 in body
    idx_343 = None
    idx_344 = None
    for idx, child in enumerate(body):
        if child.tag.endswith('p'):
            t = "".join([t_node.text for t_node in child.xpath('.//w:r//w:t') if t_node.text]).strip()
            if t == "3.4.3 Entity-Relationship Design":
                idx_343 = idx
            elif t == "3.4.4 Application Flow Design":
                idx_344 = idx

    print(f"Section 3.4.3 index in body: {idx_343}")
    print(f"Section 3.4.4 index in body: {idx_344}")
    assert idx_343 is not None and idx_344 is not None, "Failed to find 3.4.3 or 3.4.4 headings!"

    # 3. Build the two styled tables with images using python-docx
    # Table A: Peter Chen Conceptual ERD
    tbl_a = doc.add_table(rows=1, cols=1)
    tblPr_xml = (
        f'<w:tblPr {nsdecls("w")}>'
        f'  <w:tblW w:w="9026" w:type="dxa"/>'
        f'  <w:tblBorders>'
        f'    <w:top w:val="single" w:sz="4" w:space="0" w:color="auto"/>'
        f'    <w:left w:val="single" w:sz="4" w:space="0" w:color="auto"/>'
        f'    <w:bottom w:val="single" w:sz="4" w:space="0" w:color="auto"/>'
        f'    <w:right w:val="single" w:sz="4" w:space="0" w:color="auto"/>'
        f'    <w:insideH w:val="single" w:sz="4" w:space="0" w:color="auto"/>'
        f'    <w:insideV w:val="single" w:sz="4" w:space="0" w:color="auto"/>'
        f'  </w:tblBorders>'
        f'  <w:tblCellMar>'
        f'    <w:left w:w="10" w:type="dxa"/>'
        f'    <w:right w:w="10" w:type="dxa"/>'
        f'  </w:tblCellMar>'
        f'  <w:tblLook w:val="0000" w:firstRow="0" w:lastRow="0" w:firstColumn="0" w:lastColumn="0" w:noHBand="0" w:noVBand="0"/>'
        f'</w:tblPr>'
    )
    tbl_a._element.replace(tbl_a._element.tblPr, parse_xml(tblPr_xml))
    cell_a = tbl_a.rows[0].cells[0]
    tcPr_xml = (
        f'<w:tcPr {nsdecls("w")}>'
        f'  <w:tcW w:w="9026" w:type="dxa"/>'
        f'  <w:tcBorders>'
        f'    <w:top w:val="dashed" w:sz="2" w:space="0" w:color="FF6600"/>'
        f'    <w:left w:val="dashed" w:sz="2" w:space="0" w:color="FF6600"/>'
        f'    <w:bottom w:val="dashed" w:sz="2" w:space="0" w:color="FF6600"/>'
        f'    <w:right w:val="dashed" w:sz="2" w:space="0" w:color="FF6600"/>'
        f'  </w:tcBorders>'
        f'  <w:shd w:val="clear" w:color="auto" w:fill="FFF8F4"/>'
        f'  <w:tcMar>'
        f'    <w:top w:w="280" w:type="dxa"/>'
        f'    <w:left w:w="200" w:type="dxa"/>'
        f'    <w:bottom w:w="280" w:type="dxa"/>'
        f'    <w:right w:w="200" w:type="dxa"/>'
        f'  </w:tcMar>'
        f'</w:tcPr>'
    )
    cell_a._element.replace(cell_a._element.tcPr, parse_xml(tcPr_xml))
    p_a = cell_a.paragraphs[0]
    p_a.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_a = p_a.add_run()
    run_a.add_picture('docs/erd_chen.png', width=Inches(6.2))
    tbl_a_elem = tbl_a._element

    # Table B: Crow's Foot Physical Relational Schema
    tbl_b = doc.add_table(rows=1, cols=1)
    tbl_b._element.replace(tbl_b._element.tblPr, parse_xml(tblPr_xml))
    cell_b = tbl_b.rows[0].cells[0]
    cell_b._element.replace(cell_b._element.tcPr, parse_xml(tcPr_xml))
    p_b = cell_b.paragraphs[0]
    p_b.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_b = p_b.add_run()
    run_b.add_picture('docs/erd_crows_foot.png', width=Inches(6.2))
    tbl_b_elem = tbl_b._element

    # Remove tbl_a and tbl_b from the end of body so we can insert them into Section 3.4.3
    body.remove(tbl_a_elem)
    body.remove(tbl_b_elem)

    # 4. Helper function to make standard text paragraphs with proper escaping
    def make_body_p(text, bold_prefix="", space_before=0, space_after=120, jc="both"):
        xml = (
            f'<w:p {nsdecls("w")}>'
            f'  <w:pPr>'
            f'    <w:spacing w:before="{space_before}" w:after="{space_after}" w:line="360" w:lineRule="auto"/>'
            f'    <w:jc w:val="{jc}"/>'
            f'  </w:pPr>'
        )
        if bold_prefix:
            xml += (
                f'  <w:r>'
                f'    <w:rPr><w:b/><w:color w:val="000000"/></w:rPr>'
                f'    <w:t xml:space="preserve">{escape(bold_prefix)}</w:t>'
                f'  </w:r>'
            )
        if text:
            xml += (
                f'  <w:r>'
                f'    <w:rPr><w:color w:val="000000"/></w:rPr>'
                f'    <w:t xml:space="preserve">{escape(text)}</w:t>'
                f'  </w:r>'
            )
        xml += f'</w:p>'
        return parse_xml(xml)

    def make_caption_p(caption_text):
        xml = (
            f'<w:p {nsdecls("w")}>'
            f'  <w:pPr>'
            f'    <w:spacing w:before="120" w:after="160" w:line="360" w:lineRule="auto"/>'
            f'    <w:jc w:val="center"/>'
            f'  </w:pPr>'
            f'  <w:r>'
            f'    <w:rPr>'
            f'      <w:i/>'
            f'      <w:sz w:val="22"/>'
            f'      <w:color w:val="000000"/>'
            f'    </w:rPr>'
            f'    <w:t>{escape(caption_text)}</w:t>'
            f'  </w:r>'
            f'</w:p>'
        )
        return parse_xml(xml)

    def make_bullet_p(bold_title, description):
        xml = (
            f'<w:p {nsdecls("w")}>'
            f'  <w:pPr>'
            f'    <w:pStyle w:val="ListParagraph"/>'
            f'    <w:numPr>'
            f'      <w:ilvl w:val="0"/>'
            f'      <w:numId w:val="3"/>'
            f'    </w:numPr>'
            f'    <w:spacing w:before="80" w:after="80" w:line="360" w:lineRule="auto"/>'
            f'    <w:jc w:val="both"/>'
            f'  </w:pPr>'
            f'  <w:r>'
            f'    <w:rPr><w:b/><w:color w:val="000000"/></w:rPr>'
            f'    <w:t xml:space="preserve">{escape(bold_title)}: </w:t>'
            f'  </w:r>'
            f'  <w:r>'
            f'    <w:rPr><w:color w:val="000000"/></w:rPr>'
            f'    <w:t xml:space="preserve">{escape(description)}</w:t>'
            f'  </w:r>'
            f'</w:p>'
        )
        return parse_xml(xml)

    # 5. Prepare all replacement elements for Section 3.4.3
    elements_to_insert = []

    # Paragraph 1: Conceptual vs Physical Overview
    elements_to_insert.append(make_body_p(
        "To establish a comprehensive database specification that satisfies academic database design principles and serves as an exact blueprint for MySQL implementation, the system data architecture is presented at two complementary levels: a Conceptual Entity-Relationship Model (Peter Chen Notation) and a Physical Relational Database Schema (Crow's Foot Notation). The conceptual model formalizes domain entities, business semantics, and operational relationships, while the physical schema specifies database tables, exact column data types, foreign key constraints, nullability, and indexing in MySQL 8 (tipper_truck_db).",
        space_before=120, space_after=140
    ))

    # Subsection Narrative: Conceptual Model
    elements_to_insert.append(make_body_p(
        "Figure 3.5a illustrates the conceptual architecture using Peter Chen notation. Standard entity sets (USER, SAND_TYPE, TRUCK_TYPE, DELIVERY_ZONE, ORDER, ISSUE, SAND_TRUCK_PRICE, and CHATBOT_UNMATCHED_LOG) are represented as strong entities (single rectangles). The ORDER_STATUS_LOG is modeled as a weak entity (double rectangle) because its lifecycle existence and identification depend entirely on an active parent ORDER via the identifying relationship 'Has Log' (double diamond). Relationships between entities are represented as diamonds with explicit participation cardinalities (1:1, 1:N, and M:N resolved through associative pricing), and primary key attributes are underlined.",
        bold_prefix="Conceptual Entity-Relationship Model (Peter Chen Notation): ",
        space_before=100, space_after=140
    ))

    # Figure 3.5a Table Container
    elements_to_insert.append(tbl_a_elem)

    # Figure 3.5a Caption
    elements_to_insert.append(make_caption_p("Figure 3.5a Conceptual Entity-Relationship Model (Peter Chen Notation)"))

    # Subsection Narrative: Physical Relational Schema
    elements_to_insert.append(make_body_p(
        "Figure 3.5b presents the physical relational database schema in Information Engineering (Crow's Foot) notation for MySQL 8 (tipper_truck_db). The schema implements nine application tables normalized to Third Normal Form (3NF) to eliminate update anomalies while preserving referential integrity. The physical schema addresses several key real-world architectural requirements: (1) Nullable Optional Foreign Keys: Four foreign key relationships are explicitly optional (|o), including orders.assigned_operator_id (null upon order booking, assigned when dispatched), issues.order_id (null for general system inquiries), order_status_log.changed_by (null for automated system events), and chatbot_unmatched_logs.user_id (null for unauthenticated guests); (2) Financial Price Snapshotting: To prevent historical revenue audit distortion caused by subsequent catalogue price updates, the orders table snapshots price_ghs, delivery_fee_ghs, and total_ghs at the moment of order placement; (3) Dynamic Matrix Pricing: Base sand rates are decoupled into a 3×3 sand_truck_prices table with regional delivery surcharges applied from delivery_zones matching orders.region; and (4) Lifecycle Audit Logging: The order_status_log table provides an append-only audit trail of every status change (Confirmed → On The Way → Delivered).",
        bold_prefix="Physical Relational Database Schema (Crow's Foot Notation): ",
        space_before=140, space_after=140
    ))

    # Figure 3.5b Table Container
    elements_to_insert.append(tbl_b_elem)

    # Figure 3.5b Caption
    elements_to_insert.append(make_caption_p("Figure 3.5b Physical Relational Database Schema (Crow's Foot Notation)"))

    # Entity Relationships Intro
    elements_to_insert.append(make_body_p(
        "The relational foreign-key integrity constraints and participation rules across the nine domain tables are defined as follows:",
        bold_prefix="Entity Relationships and Integrity Constraints: ",
        space_before=140, space_after=100
    ))

    # Bullet points (11 rigorous bullets)
    bullets = [
        ("USERS to ORDERS (Customer)", "One-to-many (1:N). A registered client user can place multiple sand delivery orders over time; each order belongs to exactly one placing customer (orders.user_id FK, ON DELETE RESTRICT)."),
        ("USERS to ORDERS (Assigned Driver)", "One-to-many optional (1:0..N). A truck operator can be assigned to deliver zero, one, or multiple active orders; an order optionally references an assigned driver (orders.assigned_operator_id FK nullable, ON DELETE SET NULL)."),
        ("USERS to ISSUES", "One-to-many (1:N). A registered user can lodge multiple support complaints or feedback reports; each issue is owned by exactly one user (issues.user_id FK, ON DELETE CASCADE)."),
        ("USERS to ORDER_STATUS_LOG", "One-to-many optional (1:0..N). An administrator or driver can record multiple order status transitions; a status log entry optionally records the user who performed the transition (order_status_log.changed_by FK nullable)."),
        ("USERS to CHATBOT_UNMATCHED_LOGS", "One-to-many optional (1:0..N). Low-confidence chatbot inquiries optionally record the submitting user ID when authenticated (chatbot_unmatched_logs.user_id FK nullable)."),
        ("ORDERS to SAND_TYPES", "Many-to-one (N:1). Multiple delivery orders can reference the same sand catalogue item; each order has exactly one sand type (orders.sand_type_id FK, ON DELETE RESTRICT)."),
        ("ORDERS to TRUCK_TYPES", "Many-to-one (N:1). Multiple delivery orders can request the same truck capacity; each order has exactly one truck type (orders.truck_type_id FK, ON DELETE RESTRICT)."),
        ("ORDERS to ORDER_STATUS_LOG", "One-to-many identifying (1:N). An order generates a sequence of status transitions (Confirmed → On The Way → Delivered); each log record is strictly tied to one order (order_status_log.order_id FK, ON DELETE CASCADE)."),
        ("ORDERS to ISSUES", "One-to-many optional (1:0..N). An order may optionally have zero, one, or multiple associated support tickets; an issue may optionally reference a specific order (issues.order_id FK nullable, ON DELETE SET NULL)."),
        ("SAND_TYPES & TRUCK_TYPES to SAND_TRUCK_PRICES", "Composite one-to-many (1:N). Forms the 3×3 base price lookup matrix (9 unique pairs) enforced by a composite unique index on (sand_type_id, truck_type_id)."),
        ("DELIVERY_ZONES to ORDERS", "One-to-many regional lookup (1:N). The active delivery zone determines the regional surcharge applied to orders.delivery_fee_ghs matching orders.region.")
    ]

    for b_title, b_desc in bullets:
        elements_to_insert.append(make_bullet_p(b_title, b_desc))

    # Spacing before 3.4.4
    elements_to_insert.append(parse_xml(f'<w:p {nsdecls("w")}><w:pPr><w:spacing w:line="360" w:lineRule="auto"/></w:pPr></w:p>'))

    # 6. Delete old elements between idx_343 and idx_344
    elems_to_remove_count = idx_344 - (idx_343 + 1)
    print(f"Removing {elems_to_remove_count} old elements between 3.4.3 and 3.4.4...")
    for _ in range(elems_to_remove_count):
        elem = body[idx_343 + 1]
        body.remove(elem)

    # 7. Insert new elements right after idx_343
    insert_pos = idx_343 + 1
    for new_elem in elements_to_insert:
        body.insert(insert_pos, new_elem)
        insert_pos += 1

    print("All new elements successfully spliced into body!")

    # 8. Save updated document to output_file
    doc.save(output_file)
    print(f"Successfully saved updated proposal to: {output_file}")
    
    # Try copying to src_file
    try:
        shutil.copy2(output_file, src_file)
        print(f"Successfully copied updated file to: {src_file}")
    except Exception as e:
        print(f"Note: Could not overwrite {src_file} directly ({e}), but {output_file} is ready!")

if __name__ == '__main__':
    create_proposal_with_both_erds()
