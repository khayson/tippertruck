"""Fill FOCIS Appendix A–C in TipperTruck.docx for lecture review / submission."""

from __future__ import annotations

from pathlib import Path

from PIL import Image, ImageChops

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.text.paragraph import Paragraph

DOCS = Path(__file__).resolve().parent
DOCX = DOCS / "TipperTruck.docx"
WF = DOCS / "uml" / "wireframes"


LISTING_A1 = r"""<?php
declare(strict_types=1);

namespace App\Services;

class OrderStatusService
{
    private const TRANSITIONS = [
        'confirmed'  => ['on_the_way', 'cancelled'],
        'on_the_way' => ['delivered', 'cancelled'],
        'delivered'  => [],
        'cancelled'  => [],
    ];

    public function transition(Order $order, OrderStatus $to, ?User $changedBy = null, ?string $note = null): Order
    {
        return DB::transaction(function () use ($order, $to, $changedBy, $note) {
            $locked = Order::lockForUpdate()->findOrFail($order->id);
            $from = $locked->status;
            $allowed = self::TRANSITIONS[$from->value] ?? [];

            if (! in_array($to->value, $allowed, true)) {
                throw ValidationException::withMessages([
                    'status' => ["Cannot transition from {$from->label()} to {$to->label()}."],
                ]);
            }

            $locked->status = $to;
            match ($to) {
                OrderStatus::OnTheWay  => $locked->dispatched_at = now(),
                OrderStatus::Delivered => $locked->delivered_at = now(),
                default => null,
            };
            $locked->save();

            OrderStatusLog::create([
                'order_id'   => $locked->id,
                'old_status' => $from,
                'new_status' => $to,
                'changed_by' => $changedBy?->id,
                'note'       => $note,
            ]);

            return $locked->fresh();
        });
    }
}"""

LISTING_A2 = r"""<?php
declare(strict_types=1);

namespace App\Services;

class PricingService
{
    public function quote(SandType $sandType, TruckType $truckType, string $region): array
    {
        $matrixPrice = SandTruckPrice::where('sand_type_id', $sandType->id)
            ->where('truck_type_id', $truckType->id)
            ->first();

        if ($matrixPrice === null) {
            throw ValidationException::withMessages([
                'sand_type_id' => ['No price configured for this sand type and truck combination.'],
            ]);
        }

        $zone = DeliveryZone::where('region', $region)->where('is_active', true)->first();
        if ($zone === null) {
            throw ValidationException::withMessages([
                'region' => ['We do not currently deliver to this region.'],
            ]);
        }

        $base = (float) $matrixPrice->price_ghs;
        $surcharge = (float) $zone->surcharge_ghs;

        return [
            'price_ghs'        => number_format($base, 2, '.', ''),
            'delivery_fee_ghs' => number_format($surcharge, 2, '.', ''),
            'total_ghs'        => number_format($base + $surcharge, 2, '.', ''),
        ];
    }
}"""

LISTING_A3 = r"""Route::prefix('v1')->group(function () {
    Route::post('auth/register', ...)->middleware('throttle:5,1');
    Route::post('auth/login',    ...)->middleware('throttle:login');
    Route::post('auth/social',   ...)->middleware('throttle:social');
    Route::get('config', ConfigController::class);

    Route::middleware(['auth:sanctum', 'throttle:60,1'])->group(function () {
        Route::post('auth/logout', ...);
        Route::get('auth/me', ...);
        Route::post('orders', ...);
        Route::get('orders', ...);
        Route::get('orders/{order}', ...);
        Route::post('orders/{order}/cancel', ...);
        Route::post('issues', ...);
        Route::get('issues', ...);
        Route::post('chatbot/message', ...);

        Route::prefix('operator')->middleware(EnsureOperator::class)->group(function () {
            Route::get('orders', ...);
            Route::get('orders/{order}', ...);
            Route::post('orders/{order}/dispatch', ...);
            Route::post('orders/{order}/deliver', ...);
        });
    });
});"""

LISTING_A4 = r"""_dio.interceptors.add(
  InterceptorsWrapper(
    onRequest: (options, handler) async {
      final token = await _storage.read(key: _tokenKey);
      if (token != null) {
        options.headers['Authorization'] = 'Bearer $token';
      }
      handler.next(options);
    },
    onError: (error, handler) async {
      if (error.response?.statusCode == 401) {
        await clearToken();
        onUnauthorized?.call();
      }
      handler.next(error);
    },
  ),
);"""

LISTING_A5 = r"""test('order snapshots the matrix price at creation time', function () {
    $user = User::factory()->create();
    $sand = SandType::first();
    $truck = TruckType::first();
    $matrixPrice = SandTruckPrice::where('sand_type_id', $sand->id)
        ->where('truck_type_id', $truck->id)
        ->first();
    $expectedPrice = $matrixPrice->price_ghs;

    $response = $this->actingAs($user)->postJson('/api/v1/orders', validOrderPayload([
        'sand_type_id' => $sand->id,
        'truck_type_id' => $truck->id,
    ]));

    $response->assertStatus(201);
    expect($response->json('data.order.price_ghs'))->toBe($expectedPrice);

    $matrixPrice->update(['price_ghs' => 99999.00]);
    $order = Order::first();
    expect($order->price_ghs)->toBe($expectedPrice);
});

test('client-supplied price and total are rejected', function () {
    $response = $this->actingAs(User::factory()->create())
        ->postJson('/api/v1/orders', validOrderPayload([
            'price_ghs' => '999.00',
            'total_ghs' => '999.00',
        ]));
    $response->assertStatus(422)
        ->assertJsonValidationErrors(['price_ghs', 'total_ghs']);
});"""

LISTING_A6 = r"""enum OrderStatus: string
{
    case Confirmed = 'confirmed';
    case OnTheWay  = 'on_the_way';
    case Delivered = 'delivered';
    case Cancelled = 'cancelled';

    public function label(): string
    {
        return match ($this) {
            self::Confirmed => 'Confirmed',
            self::OnTheWay  => 'On The Way',
            self::Delivered => 'Delivered',
            self::Cancelled => 'Cancelled',
        };
    }

    public function progressPercent(): int
    {
        return match ($this) {
            self::Confirmed => 33,
            self::OnTheWay  => 66,
            self::Delivered => 100,
            self::Cancelled => 0,
        };
    }
}"""


class Cursor:
    def __init__(self, doc: docx.Document, ref: Paragraph):
        self.doc = doc
        self.el = ref._element
        self.parent = ref._parent

    def _insert(self, style_name: str) -> Paragraph:
        """Blank paragraph — do not clone body text (that copies keepLines / indents)."""
        new_el = OxmlElement("w:p")
        self.el.addnext(new_el)
        self.el = new_el
        para = Paragraph(new_el, self.parent)
        para.style = self.doc.styles[style_name]
        _strip_breaks(para)
        pf = para.paragraph_format
        pf.page_break_before = False
        pf.keep_together = False
        pf.space_before = Pt(0)
        pf.first_line_indent = Pt(0)
        return para

    def heading(self, level: int, text: str) -> Paragraph:
        style = {1: "Heading 1", 2: "Heading 2", 3: "Heading 3"}[level]
        p = self._insert(style)
        _set_text(p, text)
        p.paragraph_format.page_break_before = False
        p.paragraph_format.keep_together = False
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_before = Pt(8 if level == 2 else 6)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        return p

    def body(self, text: str) -> Paragraph:
        p = self._insert("Normal")
        _set_text(p, text)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.first_line_indent = Pt(0)
        p.paragraph_format.keep_together = False
        return p

    def caption(self, text: str) -> Paragraph:
        p = self._insert("Normal")
        _set_text(p, text)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.keep_together = False
        for run in p.runs:
            run.italic = True
            run.font.size = Pt(10)
        return p

    def listing_caption(self, text: str) -> Paragraph:
        p = self._insert("Normal")
        _set_text(p, text)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.keep_together = False
        p.paragraph_format.keep_with_next = True
        for run in p.runs:
            run.bold = True
            run.font.size = Pt(11)
        return p

    def code(self, text: str) -> None:
        """One shaded paragraph per line so Word can page-break inside a listing."""
        lines = [line.rstrip() for line in text.strip().splitlines()]
        compact_lines: list[str] = []
        blank = False
        for line in lines:
            if line == "":
                if blank:
                    continue
                blank = True
            else:
                blank = False
            compact_lines.append(line)
        for i, line in enumerate(compact_lines):
            p = self._insert("Normal")
            display = line.replace("\t", "    ") if line else " "
            _set_text(p, display)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(4 if i == len(compact_lines) - 1 else 0)
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.keep_together = False
            p.paragraph_format.keep_with_next = False
            for run in p.runs:
                run.font.name = "Consolas"
                run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
                run.font.size = Pt(8)
                run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)
            _shade(p, "F5F5F5")

    def figure_grid(self, items: list[tuple[Path, str]], cols: int = 2, width_in: float = 2.05) -> None:
        """Phone frames in a grid. Rows may split across pages so leftover white is small."""
        n_rows = (len(items) + cols - 1) // cols
        table = self.doc.add_table(rows=n_rows, cols=cols)
        _set_table_borders(table, visible=False)
        _set_cell_margins(table, twips=40)
        for i, (path, cap) in enumerate(items):
            r, c = divmod(i, cols)
            cell = table.cell(r, c)
            cell.text = ""
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.keep_together = False
            run = p.add_run()
            run.add_picture(str(_trim_png(path)), width=Inches(width_in))
            cap_p = cell.add_paragraph(cap)
            cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap_p.paragraph_format.space_before = Pt(1)
            cap_p.paragraph_format.space_after = Pt(3)
            cap_p.paragraph_format.line_spacing = 1.0
            cap_p.paragraph_format.keep_together = False
            for cap_run in cap_p.runs:
                cap_run.italic = True
                cap_run.font.size = Pt(9)
        self.el.addnext(table._tbl)
        self.el = table._tbl

    def table(self, headers: list[str], rows: list[list[str]]) -> None:
        table = self.doc.add_table(rows=1 + len(rows), cols=len(headers))
        if self.doc.tables:
            table.style = self.doc.tables[0].style
        _set_cell_margins(table, twips=60)
        for i, h in enumerate(headers):
            table.rows[0].cells[i].text = h
            for p in table.rows[0].cells[i].paragraphs:
                for run in p.runs:
                    run.bold = True
        for r, row in enumerate(rows, start=1):
            for c, val in enumerate(row):
                table.rows[r].cells[c].text = val
        _set_table_borders(table, visible=True)
        self.el.addnext(table._tbl)
        self.el = table._tbl


def _set_text(paragraph: Paragraph, text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def _strip_breaks(paragraph: Paragraph) -> None:
    pPr = paragraph._p.get_or_add_pPr()
    for tag in ("keepLines", "pageBreakBefore", "keepNext"):
        el = pPr.find(qn(f"w:{tag}"))
        if el is not None:
            pPr.remove(el)


def _trim_png(path: Path) -> Path:
    """Crop near-white padding so the phone frame fills the figure box."""
    out_dir = DOCS / "_crops" / "wireframes"
    out_dir.mkdir(parents=True, exist_ok=True)
    out = out_dir / path.name
    im = Image.open(path).convert("RGB")
    bg = Image.new("RGB", im.size, (255, 255, 255))
    diff = ImageChops.difference(im, bg)
    bbox = diff.getbbox()
    if bbox:
        im = im.crop(bbox)
    im.save(out, optimize=True)
    return out


def _set_cell_margins(table, twips: int = 40) -> None:
    tblPr = table._tbl.tblPr
    if tblPr is None:
        tblPr = OxmlElement("w:tblPr")
        table._tbl.insert(0, tblPr)
    for child in list(tblPr):
        if child.tag == qn("w:tblCellMar"):
            tblPr.remove(child)
    mar = OxmlElement("w:tblCellMar")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:w"), str(twips))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tblPr.append(mar)


def _shade(paragraph: Paragraph, fill: str) -> None:
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    pPr.append(shd)


def _set_table_borders(table, visible: bool) -> None:
    tbl = table._tbl
    tblPr = tbl.tblPr
    if tblPr is None:
        tblPr = OxmlElement("w:tblPr")
        tbl.insert(0, tblPr)
    for child in list(tblPr):
        if child.tag == qn("w:tblBorders"):
            tblPr.remove(child)
    borders = OxmlElement("w:tblBorders")
    val = "single" if visible else "nil"
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), val)
        el.set(qn("w:sz"), "4" if visible else "0")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    tblPr.append(borders)


def _clear_after(paragraph: Paragraph) -> None:
    """Remove every body element after this paragraph (rebuild appendices in place)."""
    el = paragraph._element.getnext()
    parent = paragraph._element.getparent()
    while el is not None:
        nxt = el.getnext()
        if el.tag == qn("w:sectPr"):
            break
        parent.remove(el)
        el = nxt


def generate(docx_path: Path = DOCX) -> Path:
    doc = docx.Document(str(docx_path))
    paras = list(doc.paragraphs)

    appendices = next(p for p in paras if p.text.strip() == "APPENDICES")
    _clear_after(appendices)

    cur = Cursor(doc, appendices)

    cur.body(
        "This section supplies the supporting material required by the Faculty of Computing "
        "and Information Systems final-year project manual. Appendix A presents selected "
        "source listings that implement the design decisions in Chapters Three and Four. "
        "Appendix B presents the implemented Flutter client screens used in the lecture "
        "review and viva demonstration. Appendix C is the user manual for the three roles "
        "that the running system supports: client, operator, and administrator."
    )

    # ----- A -----
    cur.heading(2, "Appendix A. Code Snippets")
    cur.body(
        "The listings below are excerpts from the submitted monorepo. They are included so "
        "that an examiner can inspect the status machine, server-side pricing, API surface, "
        "client networking, and the automated test that forbids client-supplied money. "
        "Unrelated imports and comments are omitted. The full files remain in api/ and mobile/."
    )

    cur.listing_caption("Listing A.1  OrderStatusService — sole writer of orders.status (api/app/Services/OrderStatusService.php)")
    cur.code(LISTING_A1)
    cur.body(
        "Listing A.1 is the implementation of the state machine specified in Chapter Three. "
        "A row lock prevents two concurrent actors from applying conflicting transitions. "
        "Filament row actions and the operator dispatch/deliver endpoints call this method. "
        "Nothing else in the codebase assigns orders.status."
    )

    cur.listing_caption("Listing A.2  PricingService — matrix price plus zone surcharge (api/app/Services/PricingService.php)")
    cur.code(LISTING_A2)
    cur.body(
        "Listing A.2 is the only place an order total is calculated. The Flutter client "
        "displays GET /config quotes and the server total_ghs. It never adds the surcharge "
        "itself. OrderService copies the three money fields onto the order row at create time."
    )

    cur.listing_caption("Listing A.3  Versioned API routes (api/routes/api.php)")
    cur.code(LISTING_A3)
    cur.body(
        "Listing A.3 is the v1 contract used by the Flutter client. Registration, login and "
        "social sign-in are throttled. Authenticated traffic is limited to sixty requests per "
        "minute. Operator routes sit behind EnsureOperator. Controllers remain thin: Form "
        "Request, service, API Resource."
    )

    cur.listing_caption("Listing A.4  ApiClient auth interceptor (mobile/lib/core/api_client.dart)")
    cur.code(LISTING_A4)
    cur.body(
        "Listing A.4 is the single Dio instance required by the client house style. The "
        "Sanctum token is read from flutter_secure_storage on every request. A 401 clears "
        "the token and returns the user to Welcome. No screen constructs its own HTTP client."
    )

    cur.listing_caption("Listing A.5  Pest: price snapshot and rejection of client money (api/tests/Feature/M3/OrderTest.php)")
    cur.code(LISTING_A5)
    cur.body(
        "Listing A.5 is one of the feature tests that GitHub Actions runs against MySQL 8. "
        "After create, raising the catalogue price does not change the stored order. Posting "
        "price_ghs or total_ghs from the client is a 422 validation error."
    )

    cur.listing_caption("Listing A.6  OrderStatus labels and progress (api/app/Enums/OrderStatus.php)")
    cur.code(LISTING_A6)
    cur.body(
        "Listing A.6 is where status_label and progress_percent originate. The Flutter "
        "tracking screen renders those fields from GET /orders/{id}. It does not keep a "
        "parallel Dart map from snake_case status to a caption or a percentage."
    )

    # ----- B -----
    cur.heading(2, "Appendix B. Screenshots")
    cur.body(
        "Figures B.1 to B.8 show the implemented Flutter client on a phone frame. "
        "Identity and address are fictional demonstration data (@tippertruck.test / "
        "0200000003). Amounts are server quotes; a live demo may show different GHS "
        "figures if the Filament matrix has been edited. The booking path is Welcome, "
        "Home, truck size, delivery, summary, payment, then tracking. Figure B.8 is "
        "Tipper Bot. Driver Home and Filament /admin are covered as procedures in "
        "Appendix C rather than extra print figures."
    )

    shots = [
        ("welcome.png", "Figure B.1  Welcome"),
        ("home.png", "Figure B.2  Home / sand types"),
        ("truck.png", "Figure B.3  Truck size"),
        ("delivery.png", "Figure B.4  Delivery location"),
        ("summary.png", "Figure B.5  Order summary"),
        ("payment.png", "Figure B.6  Payment (no PIN)"),
        ("tracking.png", "Figure B.7  Tracking (status polling)"),
        ("chatbot.png", "Figure B.8  Tipper Bot"),
    ]
    items: list[tuple[Path, str]] = []
    for fname, cap in shots:
        path = WF / fname
        if not path.exists():
            raise SystemExit(f"Missing screenshot {path}")
        items.append((path, cap))
    cur.figure_grid(items, cols=2, width_in=2.05)

    cur.body(
        "Table B.1 lists the other implemented surfaces that the lecture demonstration "
        "walks through. They follow the same Material 3 theme and the same ApiClient."
    )
    cur.listing_caption("Table B.1  Additional implemented surfaces shown in the viva")
    cur.table(
        ["Surface", "Who", "What it does"],
        [
            ["Login / Register", "Client", "Email/password or Google/Facebook; Sanctum token in secure storage"],
            ["Order history", "Client", "GET /orders with status badge; pull to refresh; offline cache"],
            ["Issue report / list", "Client", "Seven issue types; optional order link; admin response"],
            ["Profile", "Client", "Name, email, logout (revokes current token only)"],
            ["Driver Home", "Operator", "Assigned runs; Mark On The Way; Mark Delivered"],
            ["Staff portal hint", "Admin on mobile", "Directs administration to Filament /admin"],
            ["Filament /admin", "Admin", "Orders, users, issues, sand/truck types, price matrix, zones"],
        ],
    )

    # ----- C -----
    cur.heading(2, "Appendix C. User Manuals")
    cur.body(
        "This appendix is the operating manual for the submitted prototype. It assumes "
        "the API is reachable (emulator default http://10.0.2.2:8000/api/v1) and that "
        "php artisan migrate:fresh --seed has been run so the demonstration accounts exist. "
        "Passwords below are local seed values only. They must not be used in production."
    )

    cur.heading(3, "C.1 Demonstration accounts")
    cur.body(
        "Table C.1 is produced by UserSeeder. Public registration creates client accounts "
        "only. Operator and administrator accounts are created in Filament. Phone numbers "
        "are placeholders, not real Mobile Money wallets."
    )
    cur.listing_caption("Table C.1  Seeded demonstration accounts (local / viva only)")
    cur.table(
        ["Role", "Email", "Password", "Where to sign in"],
        [
            ["Client", "client@tippertruck.test", "password", "Flutter app"],
            ["Operator", "operator@tippertruck.test", "password", "Flutter app (Driver Home)"],
            ["Admin", "admin@tippertruck.test", "password", "http://<host>/admin"],
        ],
    )

    cur.heading(3, "C.2 Client: create an account and place an order")
    cur.body(
        "On Welcome, tap Get Started (register) or Login. Registration needs a full name, "
        "email, password of at least eight characters with mixed case and a number, and an "
        "optional phone. Google or Facebook may be used when GET /config lists those "
        "providers. After a successful login the app stores the Sanctum token and opens Home."
    )
    cur.body(
        "On Home, choose River Sand, Quarry Sand or Filling Sand. On the truck screen choose "
        "Small, Medium or Large. The GHS figure on each card is the matrix price for that "
        "sand and truck from GET /config. Continue to delivery. Enter recipient name, a ten-"
        "digit phone starting with 0, street, an active region (launch seed: Greater Accra "
        "or Central), city, and optional landmark and note. Save this address stores the "
        "form on the device only."
    )
    cur.body(
        "Summary shows the sand, truck, delivery block, and the server breakdown (base plus "
        "zone surcharge). Proceed to Payment. Choose Mobile Money (simulated) or cash on "
        "delivery. For MoMo, enter wallet name, phone and network (MTN, Telecel or "
        "AirtelTigo). The app does not ask for a PIN. Confirm and Place Order calls POST "
        "/orders. A 201 response opens Tracking with the new order_ref."
    )

    cur.heading(3, "C.3 Client: track, history, chatbot, and issues")
    cur.body(
        "Tracking polls GET /orders/{id} while the screen is open (eight seconds, then thirty "
        "seconds after five minutes). It stops on Delivered or Cancelled. The bar uses "
        "progress_percent from the API. A client may cancel only while the order is Confirmed. "
        "Order history lists the user's orders with a status badge. Pull down to refresh. "
        "If the radio is off, the last cached page is shown with an offline banner."
    )
    cur.body(
        "Chat opens Tipper Bot. Type a question or tap a chip. Prices in replies are read "
        "from the database at request time. If the bot is unsure it offers close topics; "
        "after repeated misses it offers to file an issue. Issues uses seven categories "
        "(late delivery, wrong sand, wrong quantity, damaged goods, payment, driver conduct, "
        "other). Link an order if the complaint is about a delivery. Admin replies appear "
        "on the same list."
    )
    cur.body(
        "Profile shows the signed-in name and email. Logout revokes only the token on this "
        "device. Other sessions remain valid until they expire (seven days) or are revoked."
    )

    cur.heading(3, "C.4 Operator: Driver Home")
    cur.body(
        "Sign in on the Flutter app as operator@tippertruck.test. The router sends role "
        "operator to Driver Home rather than the client Home. The list is GET /operator/orders "
        "and contains only orders assigned to that operator. Open a confirmed order and tap "
        "Mark On The Way (POST .../dispatch). When the load is dropped, tap Mark Delivered "
        "(POST .../deliver). Cash-on-delivery orders are marked paid when they are delivered. "
        "Pull to refresh. Logout is on the header."
    )
    cur.body(
        "Operators cannot edit prices, users, or issues. Those remain on the Filament panel. "
        "An administrator assigns the operator on the order before the driver will see it."
    )

    cur.heading(3, "C.5 Administrator: Filament panel")
    cur.body(
        "Open /admin in a browser and sign in as admin@tippertruck.test. Session auth is "
        "used here, not a Sanctum bearer token. The Orders table shows status badges and "
        "filters. Row actions Mark On The Way, Mark Delivered and Cancel call "
        "OrderStatusService; they are not raw column updates. The status log relation lists "
        "every transition. Operators who open /admin see only their assigned orders."
    )
    cur.body(
        "Users lets an admin change role (client, operator, admin) and create driver "
        "accounts. Passwords are never displayed. Sand types and truck types are full CRUD, "
        "including the public image on a sand type and the popular flag on Medium. Sand "
        "Truck Prices is the nine-cell matrix. Delivery Zones sets region and surcharge_ghs "
        "and is_active (Greater Accra GHS 0, Central GHS 400 in the seed). Issues accepts "
        "an admin_response and a resolve action."
    )
    cur.body(
        "To demonstrate tracking in a viva: place an order on the phone as the client, then "
        "on /admin (or Driver Home) mark it On The Way. Within about ten seconds the phone "
        "tracking screen should move to 66 percent without a GPS map."
    )

    cur.heading(3, "C.6 Offline behaviour and common errors")
    cur.body(
        "If the device has no network, Home can still show the last cached catalogue and "
        "order list with a banner. Book a Truck, Chat send, and Submit issue stay disabled "
        "with a short explanation rather than spinning until timeout. When the API returns "
        "422, field errors appear on the matching inputs (for example a phone that does not "
        "start with 0). 401 after token expiry returns the user to Welcome. 429 on login "
        "means too many attempts; wait for Retry-After."
    )
    cur.body(
        "A region that is not an active delivery zone is rejected at create. An inactive "
        "sand or truck type cannot be ordered. Two status clicks at the same moment cannot "
        "both succeed: the second receives 422 because of the row lock in Listing A.1."
    )

    try:
        doc.save(str(docx_path))
        print(f"Updated appendices in {docx_path}")
        return docx_path
    except PermissionError:
        alt = docx_path.with_name("TipperTruck_appendices.docx")
        doc.save(str(alt))
        print(f"DOCX_LOCKED — wrote {alt}")
        return alt


if __name__ == "__main__":
    generate()
