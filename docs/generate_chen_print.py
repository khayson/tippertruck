"""
Figure 3.5a — Chen ERD in diagrams.net / draw.io default style.

Colours and geometry follow the draw.io style palette:
  Entity        fill #DAE8FC  stroke #6C8EBF
  Relationship  fill #E1D5E7  stroke #9673A6
  Attribute     fill #FFF2CC  stroke #D6B656
  Connector     black, no arrows, orthogonal
  Font          Arial (draw.io's Helvetica on Windows)

Also writes docs/erd_chen.drawio so the figure can be opened and
edited at https://app.diagrams.net/
"""

from __future__ import annotations

from pathlib import Path
from xml.sax.saxutils import escape

from PIL import Image, ImageDraw, ImageFont

DOCS = Path(__file__).resolve().parent
OUT_PNG = DOCS / "erd_chen.png"
OUT_DRAWIO = DOCS / "erd_chen.drawio"

# 1x draw.io pixels; rendered at 2x (same as File → Export → PNG 200%)
SCALE = 2
W1, H1 = 1100, 1040
W, H = W1 * SCALE, H1 * SCALE

# draw.io default palette
BG = (255, 255, 255)
INK = (0, 0, 0)
ENTITY_FILL = (218, 232, 252)
ENTITY_STROKE = (108, 142, 191)
REL_FILL = (225, 213, 231)
REL_STROKE = (150, 115, 166)
ATTR_FILL = (255, 242, 204)
ATTR_STROKE = (214, 182, 86)
LINE = (0, 0, 0)
ATTR_LINE = (102, 102, 102)
LEGEND_MUTED = (80, 80, 80)


def font(size_1x: float, bold: bool = False) -> ImageFont.FreeTypeFont:
    name = "arialbd.ttf" if bold else "arial.ttf"
    path = Path(r"C:\Windows\Fonts") / name
    try:
        return ImageFont.truetype(str(path), round(size_1x * SCALE))
    except OSError:
        return ImageFont.load_default()


F_TITLE = font(15, True)
F_ENT = font(13, True)
F_ENT_SM = font(11, True)
F_DIA = font(11, False)
F_ATTR = font(11, False)
F_ATTR_PK = font(11, True)
F_CARD = font(12, True)
F_LEGEND = font(11, False)


def S(n: float) -> int:
    return round(n * SCALE)


def measure(draw: ImageDraw.ImageDraw, text: str, fnt) -> tuple[int, int]:
    b = draw.textbbox((0, 0), text, font=fnt)
    return b[2] - b[0], b[3] - b[1]


def text_c(draw, cx, cy, text, fnt, fill=INK):
    tw, th = measure(draw, text, fnt)
    draw.text((cx - tw / 2, cy - th / 2 - S(0.5)), text, font=fnt, fill=fill)
    return tw, th


def halo(draw, cx, cy, text, fnt, fill=INK):
    tw, th = measure(draw, text, fnt)
    x, y = cx - tw / 2, cy - th / 2 - S(0.5)
    r = S(1.5)
    for dx, dy in ((-r, 0), (r, 0), (0, -r), (0, r), (-r, -r), (r, -r), (-r, r), (r, r)):
        draw.text((x + dx, y + dy), text, font=fnt, fill=BG)
    draw.text((x, y), text, font=fnt, fill=fill)


# ------------------------------------------------------------------ layout in 1x px (10 px grid, like draw.io)
L, C, R, CH = 160, 530, 820, 990

sand = (L, 150)
zone = (C, 150)
user = (R, 150)

d_includes = (345, 260)
d_covers = (C, 260)
d_places = (675, 260)
d_initiates = (905, 260)
d_assigned = (700, 340)

d_priced = (L, 450)
order = (C, 450)
chat = (CH, 450)

d_uses = (345, 590)
d_haslog = (C, 600)
d_concerns = (675, 590)
d_reports = (R, 600)
d_changes = (710, 700)

truck = (L, 850)
log = (C, 860)
issue = (R, 850)

EW, EH = 140, 42
DW, DH = 108, 66


def generate_chen_erd(out_path: Path = OUT_PNG) -> Path:
    img = Image.new("RGB", (W, H), BG)
    draw = ImageDraw.Draw(img)
    sw = max(2, S(1.5))  # draw.io strokeWidth=1–2, doubled for 200% export
    lw = max(2, S(1.25))

    # --- small legend, no banner (draw.io users place shapes, not a title card)
    ly, lx = 28, 24
    draw.rectangle((S(lx), S(ly), S(lx + 28), S(ly + 16)), fill=ENTITY_FILL, outline=ENTITY_STROKE, width=sw)
    draw.text((S(lx + 34), S(ly + 2)), "Entity", font=F_LEGEND, fill=INK)
    dx = lx + 90
    draw.polygon(
        [(S(dx + 12), S(ly)), (S(dx + 24), S(ly + 8)), (S(dx + 12), S(ly + 16)), (S(dx), S(ly + 8))],
        fill=REL_FILL,
        outline=REL_STROKE,
        width=sw,
    )
    draw.text((S(dx + 30), S(ly + 2)), "Relationship", font=F_LEGEND, fill=INK)
    ox = lx + 230
    draw.ellipse((S(ox), S(ly), S(ox + 26), S(ly + 16)), fill=ATTR_FILL, outline=ATTR_STROKE, width=sw)
    draw.text((S(ox + 32), S(ly + 2)), "Attribute", font=F_LEGEND, fill=INK)
    px = lx + 340
    draw.ellipse((S(px), S(ly), S(px + 26), S(ly + 16)), fill=ATTR_FILL, outline=ATTR_STROKE, width=sw)
    draw.line((S(px + 4), S(ly + 13), S(px + 22), S(ly + 13)), fill=INK, width=max(1, S(1)))
    draw.text((S(px + 32), S(ly + 2)), "Key attribute", font=F_LEGEND, fill=INK)
    wx = lx + 480
    draw.rectangle((S(wx), S(ly), S(wx + 28), S(ly + 16)), fill=ENTITY_FILL, outline=ENTITY_STROKE, width=max(1, S(1)))
    draw.rectangle((S(wx + 3), S(ly + 3), S(wx + 25), S(ly + 13)), outline=ENTITY_STROKE, width=sw)
    draw.text((S(wx + 34), S(ly + 2)), "Weak entity", font=F_LEGEND, fill=INK)

    def xy(p):
        return S(p[0]), S(p[1])

    def polyline(points, dashed=False, double=False):
        pts = [xy(p) for p in points]
        if double:
            for a, b in zip(pts, pts[1:]):
                dx, dy = b[0] - a[0], b[1] - a[1]
                length = max((dx * dx + dy * dy) ** 0.5, 1)
                ox, oy = -dy / length * S(2), dx / length * S(2)
                draw.line([(a[0] + ox, a[1] + oy), (b[0] + ox, b[1] + oy)], fill=LINE, width=lw)
                draw.line([(a[0] - ox, a[1] - oy), (b[0] - ox, b[1] - oy)], fill=LINE, width=lw)
            return
        if dashed:
            dash, gap = S(7), S(5)
            for a, b in zip(pts, pts[1:]):
                x1, y1 = a
                x2, y2 = b
                dist = max(((x2 - x1) ** 2 + (y2 - y1) ** 2) ** 0.5, 1)
                step = dash + gap
                n = int(dist / step)
                for i in range(n + 1):
                    t0 = (i * step) / dist
                    t1 = min((i * step + dash) / dist, 1.0)
                    if t0 >= 1:
                        break
                    draw.line(
                        [(x1 + (x2 - x1) * t0, y1 + (y2 - y1) * t0), (x1 + (x2 - x1) * t1, y1 + (y2 - y1) * t1)],
                        fill=LINE,
                        width=lw,
                    )
            return
        draw.line(pts, fill=LINE, width=lw)

    eh, ew = EH / 2, EW / 2
    sand_b = (sand[0], sand[1] + eh)
    zone_b = (zone[0], zone[1] + eh)
    user_b = (user[0], user[1] + eh)
    user_r = (user[0] + ew, user[1])
    order_l = (order[0] - ew, order[1])
    order_r = (order[0] + ew, order[1])
    order_t = (order[0], order[1] - eh)
    order_b = (order[0], order[1] + eh)
    truck_t = (truck[0], truck[1] - eh)
    log_t = (log[0], log[1] - eh)
    issue_t = (issue[0], issue[1] - eh)
    issue_r = (issue[0] + ew, issue[1])
    chat_t = (chat[0], chat[1] - eh)
    log_r = (log[0] + 90, log[1])

    polyline([sand_b, (sand[0], d_includes[1]), d_includes])
    polyline([d_includes, (d_includes[0], order[1]), order_l])
    polyline([zone_b, d_covers], dashed=True)
    polyline([d_covers, order_t], dashed=True)
    polyline([user_b, (user[0], d_places[1]), d_places])
    polyline([d_places, (d_places[0], order[1]), order_r])
    polyline([(user[0], d_assigned[1]), d_assigned])
    polyline([d_assigned, (d_assigned[0], order[1] + 12), (order_r[0], order[1] + 12)])
    polyline([user_r, (d_initiates[0], user[1]), d_initiates])
    polyline([d_initiates, (chat[0], d_initiates[1]), chat_t])
    polyline([sand_b, d_priced])
    polyline([d_priced, truck_t])
    polyline([truck_t, (truck[0], d_uses[1]), d_uses])
    polyline([d_uses, (d_uses[0], order[1]), order_l])
    polyline([order_b, d_haslog], double=True)
    polyline([d_haslog, log_t], double=True)
    polyline([(user[0], d_reports[1]), d_reports])
    polyline([d_reports, issue_t])
    polyline([order_r, (d_concerns[0], order[1]), d_concerns])
    polyline([d_concerns, (d_concerns[0], issue[1]), issue_r])
    polyline([(user[0], d_changes[1]), d_changes])
    polyline([d_changes, (d_changes[0], log[1]), log_r])

    def draw_attr(parent, ap, name, is_pk=False):
        draw.line([xy(parent), xy(ap)], fill=ATTR_LINE, width=max(1, S(1)))
        fnt = F_ATTR_PK if is_pk else F_ATTR
        tw, th = measure(draw, name, fnt)
        rx = max(tw // 2 + S(10), S(28))
        ry = max(th // 2 + S(6), S(14))
        cx, cy = xy(ap)
        draw.ellipse((cx - rx, cy - ry, cx + rx, cy + ry), fill=ATTR_FILL, outline=ATTR_STROKE, width=sw)
        text_c(draw, cx, cy, name, fnt)
        if is_pk:
            draw.line([(cx - tw / 2, cy + th / 2 + S(1)), (cx + tw / 2, cy + th / 2 + S(1))], fill=INK, width=max(1, S(1)))

    def A(origin, dx, dy):
        return origin[0] + dx, origin[1] + dy

    draw_attr(sand, A(sand, -50, -50), "id", True)
    draw_attr(sand, A(sand, 8, -62), "name")
    draw_attr(sand, A(sand, 58, -46), "slug")

    draw_attr(zone, A(zone, -58, -50), "id", True)
    draw_attr(zone, A(zone, 0, -66), "region")
    draw_attr(zone, A(zone, 64, -48), "surcharge")

    draw_attr(user, A(user, -24, -54), "id", True)
    draw_attr(user, A(user, 30, -66), "name")
    draw_attr(user, A(user, 70, -42), "role")

    draw_attr(order, A(order, -36, -64), "id", True)
    draw_attr(order, A(order, -78, 50), "order_ref")
    draw_attr(order, A(order, 78, 50), "status")

    draw_attr(d_priced, A(d_priced, -70, -42), "price_ghs")

    draw_attr(chat, A(chat, 36, -50), "id", True)
    draw_attr(chat, A(chat, 10, 52), "top_intent")

    draw_attr(truck, A(truck, -54, 50), "id", True)
    draw_attr(truck, A(truck, 8, 64), "name")
    draw_attr(truck, A(truck, 70, 44), "capacity")

    draw_attr(log, A(log, -70, 54), "id", True)
    draw_attr(log, A(log, 8, 68), "new_status")
    draw_attr(log, A(log, 78, 48), "note")

    draw_attr(issue, A(issue, -24, 54), "id", True)
    draw_attr(issue, A(issue, 44, 64), "issue_type")
    draw_attr(issue, A(issue, 80, 36), "status")

    def draw_diamond(p, label, identifying=False):
        cx, cy = xy(p)
        rw, rh = S(DW / 2), S(DH / 2)
        pts = [(cx, cy - rh), (cx + rw, cy), (cx, cy + rh), (cx - rw, cy)]
        draw.polygon(pts, fill=REL_FILL, outline=REL_STROKE, width=sw)
        if identifying:
            rw2, rh2 = rw - S(6), rh - S(6)
            draw.polygon([(cx, cy - rh2), (cx + rw2, cy), (cx, cy + rh2), (cx - rw2, cy)], outline=REL_STROKE, width=sw)
        text_c(draw, cx, cy, label, F_DIA)

    for p, label, ident in [
        (d_includes, "Includes", False),
        (d_covers, "Covers", False),
        (d_places, "Places", False),
        (d_assigned, "Assigned", False),
        (d_initiates, "Initiates", False),
        (d_priced, "Priced At", False),
        (d_uses, "Uses", False),
        (d_haslog, "Has Log", True),
        (d_concerns, "Concerns", False),
        (d_reports, "Reports", False),
        (d_changes, "Changes", False),
    ]:
        draw_diamond(p, label, ident)

    def draw_entity(p, name, weak=False):
        cx, cy = xy(p)
        fnt = F_ENT_SM if len(name) >= 16 else F_ENT
        tw, _ = measure(draw, name, fnt)
        w = max(S(EW), tw + S(24))
        h = S(EH)
        x1, y1 = cx - w / 2, cy - h / 2
        x2, y2 = cx + w / 2, cy + h / 2
        draw.rectangle((x1, y1, x2, y2), fill=ENTITY_FILL, outline=ENTITY_STROKE, width=sw)
        if weak:
            draw.rectangle((x1 + S(5), y1 + S(5), x2 - S(5), y2 - S(5)), outline=ENTITY_STROKE, width=sw)
        text_c(draw, cx, cy, name, fnt)

    draw_entity(sand, "SAND_TYPE")
    draw_entity(zone, "DELIVERY_ZONE")
    draw_entity(user, "USER")
    draw_entity(order, "ORDER")
    draw_entity(chat, "CHATBOT_LOG")
    draw_entity(truck, "TRUCK_TYPE")
    draw_entity(log, "ORDER_STATUS_LOG", weak=True)
    draw_entity(issue, "ISSUE")

    def card(p, text, dx=0, dy=0):
        halo(draw, S(p[0] + dx), S(p[1] + dy), text, F_CARD)

    card(((sand[0] + d_includes[0]) / 2, (sand[1] + d_includes[1]) / 2), "1", 14, 0)
    card((d_includes[0], (d_includes[1] + order[1]) / 2), "N", 14, -6)
    card((zone[0], (zone[1] + d_covers[1]) / 2), "1", 16, 0)
    card((order[0], (d_covers[1] + order[1]) / 2), "N", 16, 0)
    card((user[0], (user[1] + d_places[1]) / 2), "1", -16, 0)
    card((d_places[0], d_places[1] + 28), "N", -16, 0)
    card(d_assigned, "0..1", -48, -6)
    card(d_assigned, "N", 32, 24)
    card((sand[0], (sand[1] + d_priced[1]) / 2), "M", 16, 10)
    card((truck[0], (d_priced[1] + truck[1]) / 2), "N", 16, -8)
    card((truck[0], (truck[1] + d_uses[1]) / 2), "1", 14, 0)
    card((d_uses[0], (d_uses[1] + order[1]) / 2), "N", 14, 8)
    card((order[0], (order[1] + d_haslog[1]) / 2), "1", 18, 0)
    card((log[0], (d_haslog[1] + log[1]) / 2), "N", 18, 0)
    card(d_initiates, "0..1", -8, -28)
    card((chat[0], (d_initiates[1] + chat[1]) / 2), "N", 16, -8)
    card((user[0], user[1] + 58), "1", 16, 0)
    card((issue[0], (d_reports[1] + issue[1]) / 2), "N", 16, 0)
    card(d_concerns, "0..1", -8, -24)
    card(d_concerns, "N", 22, 22)
    card(d_changes, "0..1", 44, -14)
    card(d_changes, "N", 28, 28)

    img.save(out_path, dpi=(192, 192), optimize=True)
    write_drawio(OUT_DRAWIO)
    print(f"Wrote {out_path}  {img.size[0]}x{img.size[1]}px  (draw.io 200% export style)")
    print(f"Wrote {OUT_DRAWIO}  (open in https://app.diagrams.net/)")
    return out_path


def write_drawio(path: Path) -> None:
    """Native diagrams.net file so the figure can be edited in the app."""

    cells: list[str] = []
    nid = {"n": 2}

    def next_id() -> str:
        i = nid["n"]
        nid["n"] += 1
        return str(i)

    def style_entity(weak=False) -> str:
        base = "rounded=0;whiteSpace=wrap;html=1;fillColor=#dae8fc;strokeColor=#6c8ebf;fontFamily=Helvetica;fontSize=13;fontStyle=1;fontColor=#000000;"
        return base + ("shape=ext;double=1;" if weak else "")

    def style_rel(ident=False) -> str:
        base = "rhombus;whiteSpace=wrap;html=1;fillColor=#e1d5e7;strokeColor=#9673a6;fontFamily=Helvetica;fontSize=11;fontColor=#000000;"
        return base + ("double=1;" if ident else "")

    def style_attr(pk=False) -> str:
        u = "fontStyle=4;" if pk else ""
        return f"ellipse;whiteSpace=wrap;html=1;fillColor=#fff2cc;strokeColor=#d6b656;fontFamily=Helvetica;fontSize=11;fontColor=#000000;{u}"

    def vertex(value, x, y, w, h, style) -> str:
        i = next_id()
        cells.append(
            f'<mxCell id="{i}" value="{escape(value)}" style="{style}" vertex="1" parent="1">'
            f'<mxGeometry x="{x:.0f}" y="{y:.0f}" width="{w:.0f}" height="{h:.0f}" as="geometry"/>'
            f"</mxCell>"
        )
        return i

    def edge(src, tgt, extra="") -> str:
        i = next_id()
        cells.append(
            f'<mxCell id="{i}" style="endArrow=none;html=1;rounded=0;strokeColor=#000000;strokeWidth=1;'
            f'edgeStyle=orthogonalEdgeStyle;{extra}" edge="1" parent="1" source="{src}" target="{tgt}">'
            f'<mxGeometry relative="1" as="geometry"/></mxCell>'
        )
        return i

    def label(value, x, y) -> str:
        return vertex(
            value,
            x - 16,
            y - 10,
            32,
            20,
            "text;html=1;strokeColor=none;fillColor=none;align=center;verticalAlign=middle;fontFamily=Helvetica;fontSize=12;fontStyle=1;fontColor=#000000;",
        )

    def ent(name, cx, cy, weak=False) -> str:
        return vertex(name, cx - EW / 2, cy - EH / 2, EW, EH, style_entity(weak))

    def rel(name, cx, cy, ident=False) -> str:
        return vertex(name, cx - DW / 2, cy - DH / 2, DW, DH, style_rel(ident))

    def attr(name, cx, cy, pk=False) -> str:
        w = 84 if len(name) > 6 else 60
        return vertex(name, cx - w / 2, cy - 16, w, 32, style_attr(pk))

    e_sand = ent("SAND_TYPE", *sand)
    e_zone = ent("DELIVERY_ZONE", *zone)
    e_user = ent("USER", *user)
    e_order = ent("ORDER", *order)
    e_chat = ent("CHATBOT_LOG", *chat)
    e_truck = ent("TRUCK_TYPE", *truck)
    e_log = ent("ORDER_STATUS_LOG", *log, weak=True)
    e_issue = ent("ISSUE", *issue)

    r_inc = rel("Includes", *d_includes)
    r_cov = rel("Covers", *d_covers)
    r_pla = rel("Places", *d_places)
    r_ini = rel("Initiates", *d_initiates)
    r_asg = rel("Assigned", *d_assigned)
    r_prc = rel("Priced At", *d_priced)
    r_use = rel("Uses", *d_uses)
    r_log = rel("Has Log", *d_haslog, ident=True)
    r_con = rel("Concerns", *d_concerns)
    r_rep = rel("Reports", *d_reports)
    r_chg = rel("Changes", *d_changes)

    edge(e_sand, r_inc)
    edge(r_inc, e_order)
    edge(e_zone, r_cov, extra="dashed=1;")
    edge(r_cov, e_order, extra="dashed=1;")
    edge(e_user, r_pla)
    edge(r_pla, e_order)
    edge(e_user, r_asg)
    edge(r_asg, e_order)
    edge(e_user, r_ini)
    edge(r_ini, e_chat)
    edge(e_sand, r_prc)
    edge(r_prc, e_truck)
    edge(e_truck, r_use)
    edge(r_use, e_order)
    edge(e_order, r_log)
    edge(r_log, e_log)
    edge(e_user, r_rep)
    edge(r_rep, e_issue)
    edge(e_order, r_con)
    edge(r_con, e_issue)
    edge(e_user, r_chg)
    edge(r_chg, e_log)

    pairs = [
        (e_sand, sand, (-50, -50), "id", True),
        (e_sand, sand, (8, -62), "name", False),
        (e_sand, sand, (58, -46), "slug", False),
        (e_zone, zone, (-58, -50), "id", True),
        (e_zone, zone, (0, -66), "region", False),
        (e_zone, zone, (64, -48), "surcharge", False),
        (e_user, user, (-24, -54), "id", True),
        (e_user, user, (30, -66), "name", False),
        (e_user, user, (70, -42), "role", False),
        (e_order, order, (-36, -64), "id", True),
        (e_order, order, (-78, 50), "order_ref", False),
        (e_order, order, (78, 50), "status", False),
        (r_prc, d_priced, (-70, -42), "price_ghs", False),
        (e_chat, chat, (36, -50), "id", True),
        (e_chat, chat, (10, 52), "top_intent", False),
        (e_truck, truck, (-54, 50), "id", True),
        (e_truck, truck, (8, 64), "name", False),
        (e_truck, truck, (70, 44), "capacity", False),
        (e_log, log, (-70, 54), "id", True),
        (e_log, log, (8, 68), "new_status", False),
        (e_log, log, (78, 48), "note", False),
        (e_issue, issue, (-24, 54), "id", True),
        (e_issue, issue, (44, 64), "issue_type", False),
        (e_issue, issue, (80, 36), "status", False),
    ]
    for parent, origin, off, name, pk in pairs:
        aid = attr(name, origin[0] + off[0], origin[1] + off[1], pk)
        edge(parent, aid, extra="strokeColor=#666666;")

    label("1", (sand[0] + d_includes[0]) / 2 + 14, (sand[1] + d_includes[1]) / 2)
    label("N", d_includes[0] + 14, (d_includes[1] + order[1]) / 2)
    label("1", zone[0] + 16, (zone[1] + d_covers[1]) / 2)
    label("N", order[0] + 16, (d_covers[1] + order[1]) / 2)
    label("1", user[0] - 16, (user[1] + d_places[1]) / 2)
    label("N", d_places[0] - 16, d_places[1] + 28)
    label("0..1", d_assigned[0] - 40, d_assigned[1])
    label("N", d_assigned[0] + 28, d_assigned[1] + 22)
    label("M", sand[0] + 16, (sand[1] + d_priced[1]) / 2 + 10)
    label("N", truck[0] + 16, (d_priced[1] + truck[1]) / 2)
    label("1", truck[0] + 14, (truck[1] + d_uses[1]) / 2)
    label("N", d_uses[0] + 14, (d_uses[1] + order[1]) / 2)
    label("1", order[0] + 18, (order[1] + d_haslog[1]) / 2)
    label("N", log[0] + 18, (d_haslog[1] + log[1]) / 2)
    label("0..1", d_initiates[0] - 8, d_initiates[1] - 28)
    label("N", chat[0] + 16, (d_initiates[1] + chat[1]) / 2)
    label("1", user[0] + 16, user[1] + 58)
    label("N", issue[0] + 16, (d_reports[1] + issue[1]) / 2)
    label("0..1", d_concerns[0] - 8, d_concerns[1] - 24)
    label("N", d_concerns[0] + 22, d_concerns[1] + 22)
    label("0..1", d_changes[0] + 36, d_changes[1] - 8)
    label("N", d_changes[0] + 22, d_changes[1] + 24)

    xml = f"""<mxfile host="app.diagrams.net" agent="Mozilla/5.0" version="22.1.0">
  <diagram name="Figure 3.5a Chen ERD" id="chen-erd">
    <mxGraphModel dx="1400" dy="900" grid="1" gridSize="10" guides="1" tooltips="1" connect="1" arrows="0" fold="1" page="1" pageScale="1" pageWidth="{W1}" pageHeight="{H1}" math="0" shadow="0">
      <root>
        <mxCell id="0"/>
        <mxCell id="1" parent="0"/>
        {"".join(cells)}
      </root>
    </mxGraphModel>
  </diagram>
</mxfile>
"""
    path.write_text(xml, encoding="utf-8")


def update_proposal_docx(docx_path: Path | None = None, image_path: Path = OUT_PNG) -> Path:
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt

    docx_path = docx_path or (DOCS / "TipperTruck.docx")
    doc = docx.Document(str(docx_path))

    narrative = None
    caption = None
    for p in doc.paragraphs:
        t = p.text.strip()
        if t.startswith("Conceptual Entity-Relationship Model (Peter Chen Notation):"):
            narrative = p
        if t.startswith("Figure 3.5a Conceptual Entity-Relationship Model"):
            caption = p

    if narrative is None or caption is None:
        raise SystemExit("Could not find Figure 3.5a narrative/caption in the docx.")

    for drawing in list(narrative._element.iter(qn("w:drawing"))):
        drawing.getparent().remove(drawing)

    prev = caption._element.getprevious()
    while prev is not None:
        has_blip = prev.find(".//{http://schemas.openxmlformats.org/drawingml/2006/main}blip") is not None
        text = "".join(prev.itertext()).strip()
        if has_blip and not text:
            parent = prev.getparent()
            nxt = prev.getprevious()
            parent.remove(prev)
            prev = nxt
            continue
        break

    fig = doc.add_paragraph()
    fig.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fig.paragraph_format.space_before = Pt(8)
    fig.paragraph_format.space_after = Pt(4)
    run = fig.add_run()
    run.add_picture(str(image_path), width=Inches(6.2))
    caption._element.addprevious(fig._element)

    doc.save(str(docx_path))
    print(f"Updated Figure 3.5a in {docx_path}")
    return docx_path


if __name__ == "__main__":
    generate_chen_erd()
    try:
        update_proposal_docx()
    except PermissionError:
        print("DOCX_LOCKED — close TipperTruck.docx in Word and re-run.")
