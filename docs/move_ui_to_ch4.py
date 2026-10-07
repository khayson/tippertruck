"""Move §3.4.5 screen mocks (with images) into Chapter 4 as Figures 4.1–4.8."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.text.paragraph import Paragraph

DOCX = Path(r"C:\Users\KHAYSON\Desktop\tippertruck\docs\TipperTruck.finalpdf.docx")

CAPTION_MAP = {
    "Figure 3.6 Welcome Screen Wireframe": "Figure 4.1  Welcome screen (implemented)",
    "Figure 3.7 Home / Sand Type Selection Screen Wireframe": "Figure 4.2  Home / sand type selection (implemented)",
    "Figure 3.8 Truck Type Selection Screen Wireframe": "Figure 4.3  Truck type selection (implemented)",
    "Figure 3.9 Delivery Location Form Wireframe": "Figure 4.4  Delivery location form (implemented)",
    "Figure 3.10 Order Summary Screen Wireframe": "Figure 4.5  Order summary (implemented)",
    "Figure 3.11 Payment Screen Wireframe": "Figure 4.6  Payment screen (implemented)",
    "Figure 3.12 Order Tracking Screen Wireframe": "Figure 4.7  Order tracking (implemented)",
    "Figure 3.13 Chatbot Screen Wireframe": "Figure 4.8  Chatbot / Tipper Bot (implemented)",
}

LOF_MAP = {
    6: "Figure 4.1 Welcome screen (implemented) ................................. …",
    7: "Figure 4.2 Home / sand type selection (implemented) .................... …",
    8: "Figure 4.3 Truck type selection (implemented) .......................... …",
    9: "Figure 4.4 Delivery location form (implemented) ........................ …",
    10: "Figure 4.5 Order summary (implemented) ................................. …",
    11: "Figure 4.6 Payment screen (implemented) ................................ …",
    12: "Figure 4.7 Order tracking (implemented) ................................ …",
    13: "Figure 4.8 Chatbot / Tipper Bot (implemented) .......................... …",
}


def _set_text(paragraph: Paragraph, text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def _new_paragraph_after(anchor: Paragraph, text: str, style_name: str) -> Paragraph:
    new_el = OxmlElement("w:p")
    anchor._p.addnext(new_el)
    para = Paragraph(new_el, anchor._parent)
    try:
        para.style = anchor.part.document.styles[style_name]
    except KeyError:
        pass
    _set_text(para, text)
    return para


def _new_paragraph_before(anchor: Paragraph, text: str, style_name: str) -> Paragraph:
    new_el = OxmlElement("w:p")
    anchor._p.addprevious(new_el)
    para = Paragraph(new_el, anchor._parent)
    try:
        para.style = anchor.part.document.styles[style_name]
    except KeyError:
        pass
    _set_text(para, text)
    return para


def main() -> None:
    doc = Document(str(DOCX))
    paras = list(doc.paragraphs)

    idx_345 = next(i for i, p in enumerate(paras) if p.text.strip() == "3.4.5 Interface Design")
    idx_346 = next(i for i, p in enumerate(paras) if p.text.strip() == "3.4.6 Security Considerations")
    idx_412 = next(i for i, p in enumerate(paras) if p.text.strip() == "4.1.2 Implementation Details")
    idx_413 = next(i for i, p in enumerate(paras) if p.text.strip() == "4.1.3 Testing and Quality Assurance")

    screen_start = next(
        i for i in range(idx_345, idx_346) if paras[i].text.strip().startswith("Screen 1:")
    )
    move_indices: list[int] = []
    for i in range(screen_start, idx_346):
        p = paras[i]
        t = p.text.strip()
        has_drawing = bool(p._p.findall(".//" + qn("w:drawing")))
        if t or has_drawing:
            move_indices.append(i)

    anchor_idx = idx_413 - 1
    while anchor_idx > idx_412 and not paras[anchor_idx].text.strip():
        anchor_idx -= 1
    anchor = paras[anchor_idx]

    heading = _new_paragraph_after(anchor, "Implemented user interface", "Heading 4")
    heading.paragraph_format.space_before = Pt(12)
    heading.paragraph_format.space_after = Pt(6)

    intro = _new_paragraph_after(
        heading,
        "Figures 4.1 to 4.8 show the implemented Flutter booking and assistant screens that "
        "realise the interface design in Section 3.4.5. The identity and address on the "
        "captures are fictional demonstration data (@tippertruck.test). Catalogue prices are "
        "server quotes; a live demo may show different GHS figures if the Filament matrix has "
        "been edited. The mock images are included here so Chapter Four documents the running "
        "UI; Chapter Three retains the design rationale only.",
        "Normal",
    )
    intro.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    intro.paragraph_format.space_after = Pt(8)

    cursor = intro._p
    for i in move_indices:
        el = paras[i]._element
        parent = el.getparent()
        if parent is not None:
            parent.remove(el)
        cursor.addnext(el)
        cursor = el

    # Rename captions wherever they now sit
    for p in doc.paragraphs:
        t = p.text.strip()
        if t in CAPTION_MAP:
            _set_text(p, CAPTION_MAP[t])
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.italic = True

    # Pointer under 3.4.5 (before Security)
    paras = list(doc.paragraphs)
    idx_346 = next(i for i, p in enumerate(paras) if p.text.strip() == "3.4.6 Security Considerations")
    pointer = _new_paragraph_before(
        paras[idx_346],
        "The eight primary booking and assistant screens (Welcome through Chatbot), including "
        "their mock images, are documented as the implemented Flutter UI in Chapter Four, "
        "Figures 4.1 to 4.8. Section 3.4.5 therefore states the design language only: palette, "
        "Material 3 rules, and the sequential booking path.",
        "Normal",
    )
    pointer.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    # Remove leftover empty paragraphs between 3.4.5 intro and 3.4.6
    paras = list(doc.paragraphs)
    idx_345 = next(i for i, p in enumerate(paras) if p.text.strip() == "3.4.5 Interface Design")
    idx_346 = next(i for i, p in enumerate(paras) if p.text.strip() == "3.4.6 Security Considerations")
    for i in range(idx_346 - 1, idx_345, -1):
        p = paras[i]
        t = p.text.strip()
        has_drawing = bool(p._p.findall(".//" + qn("w:drawing")))
        if "Figures 4.1 to 4.8" in t:
            continue
        if t and not t.startswith("Screen ") and not t.startswith("Figure "):
            continue
        if (not t and not has_drawing) or t.startswith("Screen ") or t.startswith("Figure 3."):
            p._element.getparent().remove(p._element)

    # List of Figures
    for p in doc.paragraphs:
        t = p.text.strip()
        for n, new in LOF_MAP.items():
            if t.startswith(f"Figure 3.{n} ") and "Wireframe" in t:
                _set_text(p, new)
                break

    # Soften 3.5 if needed
    for p in doc.paragraphs:
        if "eight-screen interface design" in p.text:
            _set_text(
                p,
                p.text.replace(
                    "eight-screen interface design, and security controls are the blueprint for implementation.",
                    "interface design language (implemented screens appear in Chapter Four), "
                    "and security controls are the blueprint for implementation.",
                ),
            )

    out = DOCX
    try:
        doc.save(str(DOCX))
        print(f"Updated {DOCX}")
    except PermissionError:
        out = DOCX.with_name("TipperTruck.finalpdf_ch4_ui.docx")
        doc.save(str(out))
        print(f"DOCX_LOCKED — wrote {out}")

    doc2 = Document(str(out))
    texts = [p.text.strip() for p in doc2.paragraphs]
    nonempty = [t for t in texts if t]
    figs4 = [t for t in nonempty if t.startswith("Figure 4.")]
    print(f"Figure 4.x body captions: {len([t for t in figs4 if '…' not in t and '....' not in t])}")
    for t in figs4:
        if "…" not in t and "...." not in t:
            print(" ", t)

    i345 = nonempty.index("3.4.5 Interface Design")
    i346 = nonempty.index("3.4.6 Security Considerations")
    between = nonempty[i345:i346]
    print("Screens left in 3.4.5:", [t for t in between if t.startswith("Screen ")])
    print("Has Ch4 UI heading:", "Implemented user interface" in nonempty)
    i_ui = nonempty.index("Implemented user interface")
    i_413 = nonempty.index("4.1.3 Testing and Quality Assurance")
    ch4_screens = [t for t in nonempty[i_ui:i_413] if t.startswith("Screen ")]
    print("Screens under Ch4 UI section:", len(ch4_screens), ch4_screens)
    drawings = sum(
        1
        for p in doc2.paragraphs
        if p._p.findall(".//" + qn("w:drawing"))
    )
    print("Total drawings in doc:", drawings)


if __name__ == "__main__":
    main()
