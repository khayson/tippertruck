"""
Figure 3.5b — Physical schema in Crow's Foot notation, diagrams.net style.

Matches draw.io Entity Relation defaults:
  Table header   fill #DAE8FC  stroke #6C8EBF
  Body rows      white / #F5F8FC
  Connectors     black orthogonal Crow's Foot (ERone / ERzeroToMany / ...)
  Font           Arial

Also writes docs/erd_crows_foot.drawio for https://app.diagrams.net/
"""

from __future__ import annotations

import math
from dataclasses import dataclass
from pathlib import Path
from xml.sax.saxutils import escape

from PIL import Image, ImageDraw, ImageFont

DOCS = Path(__file__).resolve().parent
OUT_PNG = DOCS / "erd_crows_foot.png"
OUT_DRAWIO = DOCS / "erd_crows_foot.drawio"

DPI = 400
W_IN = 6.20
H_IN = 8.15
W, H = round(W_IN * DPI), round(H_IN * DPI)

BG = (255, 255, 255)
INK = (0, 0, 0)
HEAD_FILL = (218, 232, 252)  # #DAE8FC
HEAD_STROKE = (108, 142, 191)  # #6C8EBF
ROW_A = (255, 255, 255)
ROW_B = (245, 248, 252)
PK_FILL = (108, 142, 191)
FK_FILL = (150, 115, 166)
UK_FILL = (90, 90, 90)
MUTED = (90, 90, 90)
LINE = (0, 0, 0)
LABEL_BG = (255, 255, 255)


def pt(n: float) -> int:
    return round(n * DPI / 72.0)


def stroke(n: float) -> int:
    return max(1, pt(n))


def font(size_pt: float, bold: bool = False) -> ImageFont.FreeTypeFont:
    name = "arialbd.ttf" if bold else "arial.ttf"
    path = Path(r"C:\Windows\Fonts") / name
    try:
        return ImageFont.truetype(str(path), pt(size_pt))
    except OSError:
        return ImageFont.load_default()


F_HEAD = font(9.5, True)
F_HEAD_SM = font(7.4, True)
F_NAME = font(8, False)
F_NAME_B = font(8, True)
F_NAME_SM = font(7, False)
F_NAME_SM_B = font(7, True)
F_NAME_XS = font(6.4, False)
F_NAME_XS_B = font(6.4, True)
F_TYPE = font(6.5, False)
F_KEY = font(6.5, True)
F_LABEL = font(6.8, False)
F_LEGEND = font(8, False)


def measure(draw: ImageDraw.ImageDraw, text: str, fnt) -> tuple[int, int]:
    b = draw.textbbox((0, 0), text, font=fnt)
    return b[2] - b[0], b[3] - b[1]


@dataclass
class TableBox:
    name: str
    x: int
    y: int
    w: int
    h: int
    n_rows: int
    head_h: int
    row_h: int

    def left(self, t: float = 0.5) -> tuple[int, int]:
        return self.x, int(self.y + self.h * t)

    def right(self, t: float = 0.5) -> tuple[int, int]:
        return self.x + self.w, int(self.y + self.h * t)

    def top(self, t: float = 0.5) -> tuple[int, int]:
        return int(self.x + self.w * t), self.y

    def bot(self, t: float = 0.5) -> tuple[int, int]:
        return int(self.x + self.w * t), self.y + self.h

    def row_y(self, i: int) -> int:
        return self.y + self.head_h + int((i + 0.5) * self.row_h)

    def row_left(self, i: int) -> tuple[int, int]:
        return self.x, self.row_y(i)

    def row_right(self, i: int) -> tuple[int, int]:
        return self.x + self.w, self.row_y(i)


TABLES = {
    "TRUCK_TYPES": [
        ("bigint", "id", "PK"),
        ("string", "name", ""),
        ("string", "slug", "UK"),
        ("string", "capacity_label", ""),
        ("decimal", "capacity_tonnes_min", ""),
        ("decimal", "capacity_tonnes_max", ""),
        ("decimal", "price_ghs", ""),
        ("boolean", "is_popular", ""),
        ("boolean", "is_active", ""),
        ("int", "sort_order", ""),
    ],
    "DELIVERY_ZONES": [
        ("bigint", "id", "PK"),
        ("string", "region", "UK"),
        ("decimal", "surcharge_ghs", ""),
        ("boolean", "is_active", ""),
    ],
    "SAND_TYPES": [
        ("bigint", "id", "PK"),
        ("string", "name", ""),
        ("string", "slug", "UK"),
        ("string", "description", ""),
        ("string", "icon", ""),
        ("string", "image", ""),
        ("boolean", "is_active", ""),
        ("int", "sort_order", ""),
    ],
    "USERS": [
        ("bigint", "id", "PK"),
        ("string", "name", ""),
        ("string", "email", "UK"),
        ("string", "password", ""),
        ("string", "phone", ""),
        ("string", "role", ""),
        ("string", "provider", ""),
        ("string", "provider_id", ""),
    ],
    "SAND_TRUCK_PRICES": [
        ("bigint", "id", "PK"),
        ("bigint", "sand_type_id", "FK"),
        ("bigint", "truck_type_id", "FK"),
        ("decimal", "price_ghs", ""),
    ],
    "ORDERS": [
        ("bigint", "id", "PK"),
        ("string", "order_ref", "UK"),
        ("bigint", "user_id", "FK"),
        ("bigint", "sand_type_id", "FK"),
        ("bigint", "truck_type_id", "FK"),
        ("bigint", "assigned_operator_id", "FK"),
        ("decimal", "price_ghs", ""),
        ("decimal", "delivery_fee_ghs", ""),
        ("decimal", "total_ghs", ""),
        ("string", "region", ""),
        ("string", "payment_method", ""),
        ("string", "payment_status", ""),
        ("string", "status", ""),
    ],
    "CHATBOT_UNMATCHED_LOGS": [
        ("bigint", "id", "PK"),
        ("bigint", "user_id", "FK"),
        ("text", "message", ""),
        ("string", "top_intent", ""),
        ("decimal", "confidence", ""),
    ],
    "ORDER_STATUS_LOG": [
        ("bigint", "id", "PK"),
        ("bigint", "order_id", "FK"),
        ("string", "old_status", ""),
        ("string", "new_status", ""),
        ("bigint", "changed_by", "FK"),
        ("string", "note", ""),
    ],
    "ISSUES": [
        ("bigint", "id", "PK"),
        ("bigint", "user_id", "FK"),
        ("bigint", "order_id", "FK"),
        ("string", "issue_type", ""),
        ("text", "description", ""),
        ("string", "status", ""),
        ("text", "admin_response", ""),
    ],
}

HEAD_H_PT = 15.5
ROW_H_PT = 11.0


def table_height(n_rows: int) -> int:
    return pt(HEAD_H_PT) + n_rows * pt(ROW_H_PT)


def generate_crows_erd(out_path: Path = OUT_PNG) -> Path:
    img = Image.new("RGB", (W, H), BG)
    draw = ImageDraw.Draw(img)
    lw = stroke(1.15)

    _legend(draw)

    # Two columns + a centre highway so labels and crow's feet never sit on tables.
    m_l, m_r = pt(10), pt(20)
    gw = pt(36)
    col_w = (W - m_l - m_r - gw) // 2
    x_l = m_l
    x_r = m_l + col_w + gw
    g = x_l + col_w + gw // 2
    outer = W - pt(8)

    y0 = pt(50)
    gap_sm = pt(18)
    gap_lg = pt(24)
    boxes: dict[str, TableBox] = {}

    def place(name: str, x: int, y: int, w: int) -> TableBox:
        rows = TABLES[name]
        th = table_height(len(rows))
        _draw_table(draw, x, y, w, name, rows)
        box = TableBox(name, x, y, w, th, len(rows), pt(HEAD_H_PT), pt(ROW_H_PT))
        boxes[name] = box
        return box

    t_sand = place("SAND_TYPES", x_l, y0, col_w)
    t_price = place("SAND_TRUCK_PRICES", x_l, t_sand.y + t_sand.h + gap_sm, col_w)
    t_truck = place("TRUCK_TYPES", x_l, t_price.y + t_price.h + gap_sm, col_w)
    t_zone = place("DELIVERY_ZONES", x_l, t_truck.y + t_truck.h + gap_sm, col_w)
    t_log = place("ORDER_STATUS_LOG", x_l, t_zone.y + t_zone.h + gap_lg, col_w)

    t_user = place("USERS", x_r, y0, col_w)
    t_order = place("ORDERS", x_r, t_user.y + t_user.h + gap_lg, col_w)
    t_issue = place("ISSUES", x_r, t_order.y + t_order.h + gap_lg, col_w)
    t_chat = place("CHATBOT_UNMATCHED_LOGS", x_r, t_issue.y + t_issue.h + gap_lg, col_w)

    def connect(
        pts: list[tuple[int, int]],
        start: str,
        end: str,
        label: str = "",
        dashed: bool = False,
        label_seg: int | None = None,
        nudge: tuple[int, int] = (0, 0),
    ) -> None:
        pts = _dedupe(pts)
        _polyline(draw, pts, lw, dashed)
        d_start = _dir(pts[1], pts[0])
        d_end = _dir(pts[-2], pts[-1])
        _cf_end(draw, pts[0], d_start, start)
        _cf_end(draw, pts[-1], d_end, end)
        if not label:
            return
        if label_seg is None:
            label_seg = _longest_seg(pts)
        label_seg = min(label_seg, len(pts) - 2)
        a, b = pts[label_seg], pts[label_seg + 1]
        mx = (a[0] + b[0]) // 2 + nudge[0]
        my = (a[1] + b[1]) // 2 + nudge[1]
        _edge_label(draw, mx, my, label)

    # Left column: catalogue → price matrix
    connect([t_sand.bot(0.38), t_price.top(0.38)], "1", "0..n", "RESTRICT")
    connect([t_truck.top(0.50), t_price.bot(0.50)], "1", "0..n", "RESTRICT")

    # Right column: users place / are assigned to orders
    connect(
        [t_user.bot(0.32), t_order.top(0.32)],
        "1",
        "0..n",
        "user_id  ·  RESTRICT",
        nudge=(pt(-36), 0),
    )
    connect(
        [t_user.bot(0.68), t_order.top(0.68)],
        "0..1",
        "0..n",
        "operator  ·  SET NULL",
        nudge=(pt(18), 0),
    )

    # Right column: orders → issues
    connect([t_order.bot(0.45), t_issue.top(0.45)], "0..1", "0..n", "order_id  ·  SET NULL")

    # Centre highway: SAND / TRUCK / ZONE → ORDERS
    p = t_sand.right(0.62)
    connect(
        [p, (g - pt(10), p[1]), (g - pt(10), t_order.row_y(3)), t_order.row_left(3)],
        "1",
        "0..n",
        "RESTRICT",
        label_seg=1,
        nudge=(pt(-10), 0),
    )
    p = t_truck.right(0.48)
    connect(
        [p, (g + pt(10), p[1]), (g + pt(10), t_order.row_y(4)), t_order.row_left(4)],
        "1",
        "0..n",
        "RESTRICT",
        label_seg=0,
        nudge=(pt(8), pt(-8)),
    )
    p = t_zone.row_right(1)
    connect(
        [p, (g, p[1]), (g, t_order.row_y(9)), t_order.row_left(9)],
        "1",
        "0..n",
        "no FK",
        dashed=True,
        label_seg=0,
        nudge=(0, pt(-8)),
    )

    # ORDERS → ORDER_STATUS_LOG
    p = t_order.row_left(1)
    connect(
        [p, (g + pt(10), p[1]), (g + pt(10), t_log.row_y(1)), t_log.row_right(1)],
        "1",
        "0..n",
        "CASCADE",
        label_seg=1,
        nudge=(pt(12), 0),
    )

    # USERS → ORDER_STATUS_LOG.changed_by  (down the centre highway)
    p = t_user.left(0.88)
    connect(
        [p, (g - pt(10), p[1]), (g - pt(10), t_log.row_y(4)), t_log.row_right(4)],
        "0..1",
        "0..n",
        "SET NULL",
        label_seg=2,
        nudge=(pt(-22), pt(-9)),
    )

    # Right margin: USERS → ISSUES / CHATBOT (does not cross ORDERS)
    p = t_user.right(0.28)
    connect(
        [p, (outer, p[1]), (outer, t_issue.row_y(1)), t_issue.row_right(1)],
        "1",
        "0..n",
        "",
        label_seg=2,
        nudge=(pt(-4), pt(-9)),
    )
    p = t_user.right(0.78)
    connect(
        [p, (outer - pt(11), p[1]), (outer - pt(11), t_chat.row_y(1)), t_chat.row_right(1)],
        "0..1",
        "0..n",
        "",
        label_seg=2,
        nudge=(pt(-4), pt(-9)),
    )

    bottom = max(b.y + b.h for b in boxes.values()) + pt(12)
    img = img.crop((0, 0, W, min(bottom, H)))
    img.save(out_path, dpi=(DPI, DPI), optimize=True)
    write_drawio(OUT_DRAWIO, boxes)
    img.save(DOCS / "erd.png", dpi=(DPI, DPI), optimize=True)
    img.convert("RGB").save(DOCS / "erd.jpg", quality=92, dpi=(DPI, DPI))
    print(f"Wrote {out_path}  {img.size[0]}x{img.size[1]}px  ({W_IN:.2f} x {img.size[1] / DPI:.2f} in @ {DPI} dpi)")
    print(f"Wrote {OUT_DRAWIO}")
    return out_path


def _legend(draw) -> None:
    ly = pt(7)
    draw.text((pt(10), ly), "Crow's Foot  (Information Engineering)", font=F_LEGEND, fill=INK)
    y = ly + pt(14)
    items = [
        (pt(10), "1", "one (mandatory)"),
        (pt(128), "0..1", "zero or one"),
        (pt(246), "0..n", "zero or many"),
    ]
    for x, kind, caption in items:
        x2 = x + pt(28)
        draw.line([(x, y + pt(6)), (x2, y + pt(6))], fill=LINE, width=stroke(1.1))
        _cf_end(draw, (x2, y + pt(6)), (1, 0), kind, scale=1.55)
        draw.text((x + pt(36), y), caption, font=F_LEGEND, fill=INK)
    dx0, dx1 = pt(378), pt(408)
    yy = y + pt(6)
    for i in range(0, dx1 - dx0, pt(4)):
        draw.line(
            [(dx0 + i, yy), (min(dx0 + i + pt(2.2), dx1), yy)],
            fill=LINE,
            width=stroke(1.1),
        )
    draw.text((pt(406), y), "no FK", font=F_LEGEND, fill=MUTED)


def _draw_table(draw, x, y, w, title, rows) -> None:
    head_h = pt(HEAD_H_PT)
    row_h = pt(ROW_H_PT)
    h = head_h + len(rows) * row_h
    sw = stroke(1.35)
    draw.rectangle((x, y, x + w, y + h), fill=BG, outline=HEAD_STROKE, width=sw)
    draw.rectangle((x, y, x + w, y + head_h), fill=HEAD_FILL, outline=HEAD_STROKE, width=sw)

    hf = F_HEAD_SM if len(title) > 16 else F_HEAD
    tw, th = measure(draw, title, hf)
    if tw > w - pt(8):
        hf = F_HEAD_SM
        tw, th = measure(draw, title, hf)
    draw.text((x + (w - tw) / 2, y + (head_h - th) / 2 - pt(0.3)), title, font=hf, fill=INK)

    key_pad = pt(5)
    type_gap = pt(3)
    name_x = x + pt(5)

    for i, (col_type, col_name, col_key) in enumerate(rows):
        ry = y + head_h + i * row_h
        bg = ROW_B if i % 2 else ROW_A
        draw.rectangle((x + sw, ry, x + w - sw, ry + row_h), fill=bg)

        key_w = measure(draw, col_key, F_KEY)[0] if col_key else 0
        type_w = measure(draw, col_type, F_TYPE)[0]
        key_x = x + w - key_pad - key_w if col_key else x + w - key_pad
        type_x = key_x - (type_gap + type_w if col_key else 0)
        if not col_key:
            type_x = x + w - key_pad - type_w
        name_max = type_x - name_x - pt(3)

        bold = bool(col_key)
        fonts = (
            [F_NAME_B, F_NAME_SM_B, F_NAME_XS_B] if bold else [F_NAME, F_NAME_SM, F_NAME_XS]
        )
        nfill = PK_FILL if col_key == "PK" else (FK_FILL if col_key == "FK" else INK)
        _draw_fitted(draw, col_name, (name_x, ry + pt(1.5)), name_max, fonts, nfill)
        draw.text((type_x, ry + pt(2.2)), col_type, font=F_TYPE, fill=MUTED)
        if col_key:
            kfill = PK_FILL if col_key == "PK" else (FK_FILL if col_key == "FK" else UK_FILL)
            draw.text((key_x, ry + pt(2.0)), col_key, font=F_KEY, fill=kfill)


def _draw_fitted(draw, text, xy, max_w, fonts, fill) -> None:
    x, y = xy
    for fnt in fonts:
        tw, _ = measure(draw, text, fnt)
        if tw <= max_w:
            draw.text((x, y), text, font=fnt, fill=fill)
            return
    fnt = fonts[-1]
    t = text
    while t and measure(draw, t + "…", fnt)[0] > max_w:
        t = t[:-1]
    draw.text((x, y), (t + "…") if t else text, font=fnt, fill=fill)


def _dedupe(pts: list[tuple[int, int]]) -> list[tuple[int, int]]:
    out: list[tuple[int, int]] = [pts[0]]
    for p in pts[1:]:
        if p != out[-1]:
            out.append(p)
    return out


def _longest_seg(pts: list[tuple[int, int]]) -> int:
    best_i, best_d = 0, -1.0
    for i in range(len(pts) - 1):
        d = math.hypot(pts[i + 1][0] - pts[i][0], pts[i + 1][1] - pts[i][1])
        if d > best_d:
            best_i, best_d = i, d
    return best_i


def _polyline(draw, pts, width, dashed=False):
    if dashed:
        for a, b in zip(pts, pts[1:]):
            x1, y1 = a
            x2, y2 = b
            dist = max(math.hypot(x2 - x1, y2 - y1), 1)
            dash, gap = pt(3.5), pt(2.4)
            step = dash + gap
            n = int(dist / step)
            for i in range(n + 1):
                t0 = (i * step) / dist
                t1 = min((i * step + dash) / dist, 1.0)
                if t0 >= 1:
                    break
                draw.line(
                    [
                        (x1 + (x2 - x1) * t0, y1 + (y2 - y1) * t0),
                        (x1 + (x2 - x1) * t1, y1 + (y2 - y1) * t1),
                    ],
                    fill=LINE,
                    width=width,
                )
    else:
        draw.line(pts, fill=LINE, width=width)


def _dir(a, b) -> tuple[float, float]:
    dx, dy = b[0] - a[0], b[1] - a[1]
    length = max(math.hypot(dx, dy), 1)
    return dx / length, dy / length


def _cf_end(draw, pos, into_table, kind: str, scale: float = 1.0) -> None:
    """Crow's Foot terminator sitting *outside* the table along the incoming wire."""
    x, y = pos
    ux, uy = into_table
    px, py = -uy, ux
    s = pt(7.0) * scale
    w = max(1, stroke(0.95))

    def bar(dist):
        cx = x - ux * dist
        cy = y - uy * dist
        half = s * 0.62
        draw.line([(cx - px * half, cy - py * half), (cx + px * half, cy + py * half)], fill=LINE, width=w)

    def circle(dist):
        r = max(pt(2.0), s * 0.30)
        cx = x - ux * (dist + r)
        cy = y - uy * (dist + r)
        draw.ellipse((cx - r, cy - r, cx + r, cy + r), outline=LINE, width=w, fill=BG)

    def crow(dist):
        tip_x = x - ux * dist
        tip_y = y - uy * dist
        length = s * 1.55
        inset = max(2, pt(0.7))
        for deg in (-50, 0, 50):
            rad = math.radians(deg)
            ax, ay = -ux, -uy
            rx = ax * math.cos(rad) - ay * math.sin(rad)
            ry = ax * math.sin(rad) + ay * math.cos(rad)
            draw.line(
                [
                    (tip_x + rx * inset, tip_y + ry * inset),
                    (tip_x + rx * length, tip_y + ry * length),
                ],
                fill=LINE,
                width=w,
            )

    if kind == "1":
        bar(s * 0.20)
        bar(s * 0.58)
    elif kind == "0..1":
        circle(s * 1.15)
        bar(s * 0.22)
    elif kind == "0..n":
        circle(s * 1.90)
        crow(s * 0.08)
    elif kind == "1..n":
        bar(s * 0.20)
        crow(s * 0.85)


def _edge_label(draw, x, y, text: str) -> None:
    tw, th = measure(draw, text, F_LABEL)
    pad_x, pad_y = pt(2.2), pt(1.0)
    box = (x - tw / 2 - pad_x, y - th / 2 - pad_y, x + tw / 2 + pad_x, y + th / 2 + pad_y)
    draw.rectangle(box, fill=LABEL_BG, outline=(170, 170, 170), width=stroke(0.55))
    draw.text((x - tw / 2, y - th / 2 - pt(0.2)), text, font=F_LABEL, fill=INK)


def write_drawio(path: Path, boxes: dict[str, TableBox]) -> None:
    cells: list[str] = []
    nid = {"n": 2}

    def next_id() -> str:
        i = str(nid["n"])
        nid["n"] += 1
        return i

    def swimlane(name: str, box: TableBox, rows) -> str:
        sid = next_id()
        k = 0.45
        x, y, w = box.x * k, box.y * k, box.w * k
        start = 26
        row_h = 18
        h = start + len(rows) * row_h
        style = (
            "swimlane;fontStyle=1;childLayout=stackLayout;horizontal=1;startSize=26;"
            "fillColor=#dae8fc;horizontalStack=0;resizeParent=1;resizeLast=0;"
            "collapsible=0;marginBottom=0;whiteSpace=wrap;html=1;"
            "strokeColor=#6c8ebf;fontFamily=Helvetica;fontSize=12;fontColor=#000000;"
        )
        cells.append(
            f'<mxCell id="{sid}" value="{escape(name)}" style="{style}" vertex="1" parent="1">'
            f'<mxGeometry x="{x:.0f}" y="{y:.0f}" width="{w:.0f}" height="{h}" as="geometry"/>'
            f"</mxCell>"
        )
        for col_type, col_name, col_key in rows:
            cid = next_id()
            extra = f"{col_key}  " if col_key else ""
            label = f"{extra}{col_name}  :  {col_type}"
            cells.append(
                f'<mxCell id="{cid}" value="{escape(label)}" '
                f'style="text;strokeColor=none;fillColor=none;align=left;verticalAlign=middle;'
                f'spacingLeft=4;spacingRight=4;overflow=hidden;rotatable=0;points=[[0,0.5],[1,0.5]];'
                f'portConstraint=eastwest;fontFamily=Helvetica;fontSize=11;" vertex="1" parent="{sid}">'
                f'<mxGeometry y="{start}" width="{w:.0f}" height="{row_h}" as="geometry"/>'
                f"</mxCell>"
            )
            start += row_h
        return sid

    ids = {name: swimlane(name, box, TABLES[name]) for name, box in boxes.items()}

    def edge(src, tgt, start_arrow, end_arrow, dashed=False, label=""):
        eid = next_id()
        dash = "dashed=1;" if dashed else ""
        lab = f' value="{escape(label)}"' if label else ""
        cells.append(
            f'<mxCell id="{eid}"{lab} style="edgeStyle=orthogonalEdgeStyle;rounded=0;orthogonalLoop=1;'
            f'jettySize=auto;html=1;endArrow={end_arrow};startArrow={start_arrow};endFill=0;startFill=0;'
            f'strokeColor=#000000;{dash}fontFamily=Helvetica;fontSize=10;" edge="1" parent="1" '
            f'source="{src}" target="{tgt}"><mxGeometry relative="1" as="geometry"/></mxCell>'
        )

    edge(ids["SAND_TYPES"], ids["ORDERS"], "ERone", "ERzeroToMany", label="sand_type_id RESTRICT")
    edge(ids["SAND_TYPES"], ids["SAND_TRUCK_PRICES"], "ERone", "ERzeroToMany", label="sand_type_id")
    edge(ids["TRUCK_TYPES"], ids["ORDERS"], "ERone", "ERzeroToMany", label="truck_type_id RESTRICT")
    edge(ids["TRUCK_TYPES"], ids["SAND_TRUCK_PRICES"], "ERone", "ERzeroToMany", label="truck_type_id")
    edge(ids["DELIVERY_ZONES"], ids["ORDERS"], "ERone", "ERzeroToMany", True, "region match (no FK)")
    edge(ids["USERS"], ids["ORDERS"], "ERone", "ERzeroToMany", label="user_id RESTRICT")
    edge(ids["USERS"], ids["ORDERS"], "ERzeroToOne", "ERzeroToMany", label="assigned_operator_id SET NULL")
    edge(ids["USERS"], ids["CHATBOT_UNMATCHED_LOGS"], "ERzeroToOne", "ERzeroToMany", label="user_id CASCADE")
    edge(ids["USERS"], ids["ISSUES"], "ERone", "ERzeroToMany", label="user_id CASCADE")
    edge(ids["ORDERS"], ids["ORDER_STATUS_LOG"], "ERone", "ERzeroToMany", label="order_id CASCADE")
    edge(ids["USERS"], ids["ORDER_STATUS_LOG"], "ERzeroToOne", "ERzeroToMany", label="changed_by SET NULL")
    edge(ids["ORDERS"], ids["ISSUES"], "ERzeroToOne", "ERzeroToMany", label="order_id SET NULL")

    xml = f"""<mxfile host="app.diagrams.net" agent="Mozilla/5.0" version="22.1.0">
  <diagram name="Figure 3.5b Crow's Foot" id="crows-erd">
    <mxGraphModel dx="1400" dy="900" grid="1" gridSize="10" guides="1" tooltips="1" connect="1" arrows="0" fold="1" page="1" pageScale="1" pageWidth="1100" pageHeight="1250" math="0" shadow="0">
      <root>
        <mxCell id="0"/>
        <mxCell id="1" parent="0"/>
        {"".join(cells)}
      </root>
    </mxGraphModel>
  </diagram>
</mxfile>
"""
    path.write_text(xml, encoding="utf-8")


def update_proposal_docx(docx_path: Path | None = None, image_path: Path = OUT_PNG) -> Path:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt
    import docx

    docx_path = docx_path or (DOCS / "TipperTruck.docx")
    doc = docx.Document(str(docx_path))

    caption = None
    for p in doc.paragraphs:
        if p.text.strip().startswith("Figure 3.5b Physical Relational Database Schema"):
            caption = p
            break
    if caption is None:
        raise SystemExit("Could not find Figure 3.5b caption in the docx.")

    prev = caption._element.getprevious()
    while prev is not None:
        tag = prev.tag.split("}")[-1]
        has_blip = prev.find(".//{http://schemas.openxmlformats.org/drawingml/2006/main}blip") is not None
        text = "".join(prev.itertext()).strip()
        if (tag == "tbl" and has_blip) or (has_blip and not text):
            nxt = prev.getprevious()
            prev.getparent().remove(prev)
            prev = nxt
            continue
        break

    fig = doc.add_paragraph()
    fig.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fig.paragraph_format.space_before = Pt(8)
    fig.paragraph_format.space_after = Pt(4)
    run = fig.add_run()
    run.add_picture(str(image_path), width=Inches(6.2))
    caption._element.addprevious(fig._element)

    doc.save(str(docx_path))
    print(f"Updated Figure 3.5b in {docx_path}")
    return docx_path


if __name__ == "__main__":
    generate_crows_erd()
    try:
        update_proposal_docx()
    except PermissionError:
        print("DOCX_LOCKED — close TipperTruck.docx in Word and re-run.")
